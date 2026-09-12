import re
import zipfile
import streamlit as st
import pandas as pd
import datetime
import pytz
import os
import requests
import io
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import smtplib
from email.mime.text import MIMEText
from email.header import Header

# Set page config for web app
st.set_page_config(
    page_title="園區情緒觀察營 週總結回報平台",
    page_icon="🌟",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom Styling
st.markdown("""
<style>
    /* 全域清新水彩漸層背景與精緻字型 */
    html, body, [data-testid="stAppViewContainer"], [data-testid="stHeader"] {
        background-color: #F6FAF4 !important;
        background-image: linear-gradient(135deg, #F3F8F2 0%, #FAF5EF 100%) !important;
    }
    
    /* 調整字型與主要內文顏色 (移除 span 以避免破壞 Streamlit Chevron 箭頭與排版) */
    .stApp, p, label, li {
        font-family: 'PingFang TC', 'Microsoft JhengHei', sans-serif !important;
        color: #2E4436 !important; /* 深森林綠，確保手機上文字高度可讀 */
    }

    /* 網頁大標題：清新優雅水彩風 */
    .report-title {
        font-family: 'PingFang TC', 'Microsoft JhengHei', sans-serif;
        font-size: 2.3rem;
        font-weight: 800;
        color: #3C6E47; /* 溫和的水彩葉綠色 */
        text-align: center;
        margin-top: 15px;
        margin-bottom: 8px;
        text-shadow: 1px 1px 4px rgba(60, 110, 71, 0.15);
    }
    
    .report-subtitle {
        font-family: 'PingFang TC', 'Microsoft JhengHei', sans-serif;
        font-size: 1.1rem;
        color: #7A695A; /* 柔和的泥土暖褐 */
        text-align: center;
        margin-bottom: 30px;
        letter-spacing: 0.05em;
    }
    
    /* 提示區塊：淡淡的水彩綠葉底色 + 柔和左邊框 */
    .group-box {
        background-color: #EBF3EC !important;
        color: #2E5638 !important;
        padding: 16px;
        border-radius: 12px;
        border-left: 6px solid #8FA893 !important; /* 柔軟水彩綠 */
        box-shadow: 2px 4px 15px rgba(143, 168, 147, 0.12);
        margin-bottom: 25px;
        font-size: 1.05rem;
    }
    
    /* 提交成功區塊：粉嫩花瓣綠意感 */
    .success-box {
        background-color: #F0F6F2 !important;
        color: #2A4D3B !important;
        padding: 22px;
        border-radius: 14px;
        border: 1px solid #D5E5DE !important;
        box-shadow: 3px 6px 18px rgba(47, 82, 62, 0.08);
        margin-top: 15px;
    }
    
    /* LINE 關懷與郵件範本：微風暖沙水彩調 */
    .reminder-template {
        background-color: #FCF5EC !important;
        color: #6D5144 !important;
        padding: 18px;
        border-radius: 10px;
        border-left: 5px solid #E2A784 !important; /* 柔粉橘 */
        font-family: monospace;
        white-space: pre-wrap;
        margin-top: 10px;
        box-shadow: 1px 3px 10px rgba(226, 167, 132, 0.1);
    }
    
    .email-box {
        background-color: #F1F5F9 !important;
        color: #3B4B5E !important;
        padding: 15px;
        border-radius: 10px;
        border-left: 5px solid #8EA7BA !important; /* 水彩天藍 */
        font-family: monospace;
        white-space: pre-wrap;
        margin-top: 10px;
        box-shadow: 1px 3px 10px rgba(142, 167, 186, 0.1);
    }

    /* v23 修正：調整 Streamlit 按鈕與匯出下載按鈕，使其統一符合精美水彩綠色調，避免下載按鈕呈灰色/黑色 */
    div.stButton > button, div.stDownloadButton > button {
        background-color: #8FA893 !important; /* 水彩綠 */
        color: white !important;
        border-radius: 20px !important;
        border: none !important;
        box-shadow: 1px 3px 8px rgba(143, 168, 147, 0.25) !important;
        transition: all 0.3s ease !important;
        font-weight: bold !important;
        padding: 0.5rem 1.5rem !important;
    }
    div.stButton > button:hover, div.stDownloadButton > button:hover {
        background-color: #76907B !important;
        box-shadow: 1px 4px 12px rgba(118, 144, 123, 0.35) !important;
        transform: translateY(-1px) !important;
    }

    /* 頁面展開區塊 Header 圓角與大方背景，杜絕重疊字 */
    div[data-testid="stExpander"] {
        background-color: #FFFFFF !important;
        border-radius: 12px !important;
        border: 1px solid #EAE0D5 !important;
        box-shadow: 1px 3px 10px rgba(122, 105, 90, 0.05) !important;
        margin-bottom: 15px !important;
    }
    
    div[data-testid="stExpander"] details summary {
        background-color: #FAF5EF !important;
        border-radius: 10px !important;
        color: #2E4436 !important;
        padding: 12px 18px !important;
        font-weight: bold !important;
    }

    /* 確保輸入框、文字區域、下拉選單均使用純白底、深綠字，絕無黑底，高對比度 */
    input[type="text"], input[type="password"], input[type="number"], textarea, select {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-radius: 8px !important;
        border: 1px solid #D5E5DE !important;
    }
    
    /* 下拉選單主體 */
    div[data-baseweb="select"] {
        background-color: #FFFFFF !important;
        border-radius: 8px !important;
    }
    
    /* 僅針對下拉選單內的選項與選取內容，不碰結構性 div，根除疊字 */
    div[data-baseweb="select"] div[data-testid="stMarkdownContainer"] p,
    div[data-baseweb="select"] div,
    div[data-baseweb="select"] span {
        color: #2E4436 !important;
    }
    
    /* 當焦點在輸入框或文字區時的文字與背景 */
    input:focus, textarea:focus {
        color: #2E4436 !important;
        background-color: #FFFFFF !important;
    }

    /* 下拉彈窗（Popover）本身 */
    div[data-baseweb="popover"] {
        background-color: #FFFFFF !important;
        border-radius: 10px !important;
        box-shadow: 0px 4px 15px rgba(46, 68, 54, 0.15) !important;
    }

    /* 懸浮清單與單個選項 */
    ul[role="listbox"], li[role="option"] {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
    }

    /* 確保側邊欄（Sidebar）在任何模式下都是溫暖的水彩粘土白，而非黑/深灰色 */
    section[data-testid="stSidebar"], [data-testid="stSidebarUserContent"], div[data-testid="stSidebarCollapseButton"] {
        background-color: #FAF5EF !important;
        background-image: linear-gradient(180deg, #F3F8F2 0%, #FAF5EF 100%) !important;
        color: #2E4436 !important;
    }
    section[data-testid="stSidebar"] p, 
    section[data-testid="stSidebar"] span, 
    section[data-testid="stSidebar"] label, 
    section[data-testid="stSidebar"] h1, 
    section[data-testid="stSidebar"] h2, 
    section[data-testid="stSidebar"] h3 {
        color: #2E4436 !important;
    }

    /* 確保表格 Dataframe 在任何模式下皆為白色背景與深色文字，全置中對齊防看錯行 */
    div[data-testid="stTable"] table, div[data-testid="stDataFrame"] table {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-collapse: separate !important;
        border-spacing: 0 !important;
        width: 100% !important;
    }
    div[data-testid="stTable"] th, div[data-testid="stDataFrame"] th {
        background-color: #EBF3EC !important;
        color: #1B5E20 !important;
        text-align: center !important; /* 標頭置中對齊 */
        vertical-align: middle !important;
    }
    div[data-testid="stTable"] td, div[data-testid="stDataFrame"] td {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-top: 1px solid #EAE0D5 !important; /* 補上頂底邊框線防止 separate 模式下邊框不見 */
        border-bottom: 1px solid #EAE0D5 !important;
        text-align: center !important; /* 數據置中對齊，防手機上看錯行 */
        vertical-align: middle !important;
    }

    /* 💡 v38 調整：全螢幕扁平外框，移除 max-height 與捲動限制，使其變回一般大方表格 */
    div[data-testid="stTable"] {
        overflow-x: auto !important;
        position: relative !important;
        border: 1px solid #EAE0D5 !important;
        border-radius: 12px !important;
        box-shadow: inset 0 1px 3px rgba(0,0,0,0.02) !important;
        background-color: #FFFFFF !important;
        -webkit-overflow-scrolling: touch !important;
    }

    /* v38 新增: 系統全局數據篩選器專屬的「視窗滾動容器」與「頂端列凍結（Sticky Header）」 */
    .scroll-table-container {
        max-height: 450px !important;
        overflow-y: auto !important;
        overflow-x: auto !important;
        position: relative !important;
        border: 1px solid #EAE0D5 !important;
        border-radius: 12px !important;
        box-shadow: inset 0 1px 3px rgba(0,0,0,0.02) !important;
        background-color: #FFFFFF !important;
        -webkit-overflow-scrolling: touch !important; /* 確保 iOS 手機端滑動流暢 */
        margin-top: 10px;
        margin-bottom: 20px;
    }
    .scroll-table-container table {
        width: 100% !important;
        border-collapse: separate !important;
        border-spacing: 0 !important;
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        font-size: 0.95rem;
    }
    /* 保持表頭固定在頂部 */
    .scroll-table-container thead,
    .scroll-table-container thead tr,
    .scroll-table-container th {
        position: sticky !important;
        top: 0 !important;
        z-index: 99 !important; /* 確保在滾動時絕對壓在下方資料之上 */
    }
    .scroll-table-container th {
        background-color: #EBF3EC !important;
        color: #1B5E20 !important;
        text-align: center !important;
        vertical-align: middle !important;
        padding: 10px 15px !important;
        font-weight: bold !important;
        box-shadow: 0 1px 0 #D5E5DE, 0 2px 4px rgba(0,0,0,0.04) !important; /* 加上下分割陰影 */
        border: none !important;
    }
    .scroll-table-container td {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-top: 1px solid #EAE0D5 !important;
        border-bottom: 1px solid #EAE0D5 !important;
        text-align: center !important;
        vertical-align: middle !important;
        padding: 10px 15px !important;
    }

    /* 確保 Metric 數據指標卡片也有高對比度深色文字與白/淺色底色，在任何情況下都不會呈現黑底 */
    div[data-testid="stMetricValue"] > div {
        color: #1B5E20 !important; /* 清新綠數字 */
    }
    div[data-testid="stMetricLabel"] > div {
        color: #7A695A !important; /* 泥土腳標題 */
    }
    div[data-testid="stMetric"] {
        background-color: #FAF5EF !important;
        border-radius: 12px !important;
        padding: 12px !important;
        border: 1px solid #EAE0D5 !important;
        box-shadow: 1px 3px 10px rgba(122, 105, 90, 0.08) !important;
    }


    /* v22 新增：強制程式碼區塊 (st.code) 呈現極緻高對比度淺底深綠字，杜絕黑底 */
    div[data-testid="stCodeBlock"], 
    div[data-testid="stCodeBlock"] pre, 
    div[data-testid="stCodeBlock"] code,
    div[data-testid="stCodeBlock"] span,
    code, pre {
        background-color: #FAF5EF !important;
        color: #2E4436 !important;
        border-radius: 8px !important;
        border: 1px solid #EAE0D5 !important;
    }

    /* v22 新增：強制 Multiselect 選項標籤 (Tag/Chip) 在暗黑模式下為淺底深綠字，杜絕黑底 */
    div[data-baseweb="tag"] {
        background-color: #EBF3EC !important;
        color: #2E4436 !important;
        border: 1px solid #D5E5DE !important;
    }
    div[data-baseweb="tag"] span {
        color: #2E4436 !important;
    }

    /* v22 新增：強制所有輸入、文字區域、下拉選單容器保持高對比度純白底色，在任何情況下都不會呈現黑底 */
    div[data-testid="stTextArea"] textarea, 
    div[data-testid="stTextInput"] input, 
    div[data-testid="stNumberInput"] input, 
    div[data-testid="stSelectbox"] select,
    div[data-testid="stTextArea"] div, 
    div[data-testid="stTextInput"] div, 
    div[data-testid="stNumberInput"] div, 
    div[data-testid="stSelectbox"] div,
    div[data-baseweb="input"], 
    div[data-baseweb="textarea"],
    div[data-baseweb="select"] {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
    }


    /* v23 新增：系統全局數據篩選器的兩個選單 (Multiselect) 及其內部容器強制白底深綠字，防止暗黑模式下變黑 */
    div[data-testid="stMultiSelect"], 
    div[data-testid="stMultiSelect"] > div, 
    div[data-testid="stMultiSelect"] div[data-baseweb="select"],
    div[data-testid="stMultiSelect"] div[role="combobox"],
    div[data-testid="stMultiSelect"] div[role="button"],
    div[data-testid="stMultiSelect"] input {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
    }
    /* 保持選取後的標籤 (Tag/Chip) 為淡雅水彩草綠底，使視覺對比更立體 */
    div[data-baseweb="tag"] {
        background-color: #EBF3EC !important;
        color: #2E4436 !important;
        border: 1px solid #D5E5DE !important;
    }
    div[data-baseweb="tag"] span {
        color: #2E4436 !important;
    }

</style>
""", unsafe_allow_html=True)

# Render Header
st.markdown('<div class="report-title">園區情緒觀察營 週總結回報平台</div>', unsafe_allow_html=True)


# 1~18組完整分組名單 (來源自保溫分組名單)
# 1~18組完整分組名單 (嚴格依據 2026園區情緒觀察營-學員總名冊0805.pdf 整理，並在名字前方加上組內編號)
groups_data = {
    "第 1 組": {
        "leader": "曾旭正、李麗斐（副）",
        "members": [
            "1-1 曾旭正", "1-2 李麗斐", "1-3 蔡淵輝", "1-4 薛榕婷", "1-5 關婉玲", "1-6 周文祥", "1-7 馮豐隆", "1-8 郭擺灧", "1-9 鍾昆原", "1-10 張筱雯", "1-11 王坤德", "1-12 吳美毅", "1-13 曾美芳", "1-14 陳淑芬", "1-15 蔡沛妤", "1-16 黃淑清", "1-17 陳萬和", "1-18 梁庭", "1-19 嚴惠英"
        ]
    },
    "第 2 組": {
        "leader": "楊韻蓉",
        "members": [
            "2-1 林文華", "2-2 林佳慧", "2-3 楊韻蓉", "2-4 林翔均", "2-5 溫婷伊", "2-6 陳田恬", "2-7 廖袖婷", "2-8 王均晨", "2-9 蔡宜均", "2-10 吳煥崇", "2-11 黃曉鈺", "2-12 葉桂香", "2-13 陳秀花", "2-14 劉德政", "2-15 吳振仁"
        ]
    },
    "第 3 組": {
        "leader": "陳威廷",
        "members": [
            "3-1 陳威廷", "3-2 紀辰諭", "3-3 劉俊麟", "3-4 張又方", "3-5 謝耀德", "3-6 葉迷妮", "3-7 王韶怡", "3-8 郭芳佑", "3-9 王碧宏", "3-10 柯春黛", "3-11 黃玉宙", "3-12 何進德", "3-13 蕭名權", "3-14 戈德", "3-15 蘇献珍", "3-16 張顥嚴", "3-17 陳淑媛", "3-18 林秀婷", "3-19 柯春僖", "3-20 簡麗分", "3-21 賴彩鈴", "3-22 林耘庄", "3-23 黃如榛", "3-24 許媛婷", "3-25 陳薇莉", "3-26 張開明", "3-27 吳翊菱", "3-28 鍾亞昭", "3-29 林金泉"
        ]
    },
    "第 4 組": {
        "leader": "王國美",
        "members": [
            "4-1 王國美", "4-2 黃柏豪", "4-3 謝佳容", "4-4 劉以婕", "4-5 留千惠", "4-6 張玉芳", "4-7 吳念真", "4-8 陳芸詞", "4-9 陳揚旻", "4-10 鄭翊汝", "4-11 李濱如", "4-12 洪筱婷", "4-13 張麗卿", "4-14 陳聖昀", "4-15 曾麗珍", "4-16 賀正楨", "4-17 潘秋華", "4-18 鄭美華", "4-19 羅怡和", "4-20 周秀芬", "4-21 侯得正"
        ]
    },
    "第 5 組": {
        "leader": "郭宗川",
        "members": [
            "5-1 郭宗川", "5-2 劉馨檀", "5-3 黃瓊慧", "5-4 石俊傑", "5-5 張志豪", "5-6 林淑寧", "5-7 林麗雪", "5-8 陳毓旻", "5-9 楊舒涵", "5-10 李平裕", "5-11 李俊男", "5-12 尤躍勳", "5-13 戴世昌", "5-14 黃壬癸", "5-15 蔡宗佑", "5-16 邱淑春", "5-17 郭富貴", "5-18 李麗卿", "5-19 蕭玉卿", "5-20 陳秀美", "5-21 劉玉如", "5-22 賴幸絹"
        ]
    },
    "第 6 組": {
        "leader": "林家敬",
        "members": [
            "6-1 林家敬", "6-2 劉洹岑", "6-3 林德慶", "6-4 彭聖森", "6-5 粘智淳", "6-6 羅賢芳", "6-7 陳君妮", "6-8 徐立昇", "6-9 粘振清", "6-10 劉潓雪", "6-11 潘玟玲", "6-12 韓倩如", "6-13 黃乙娟", "6-14 廖鶴翔", "6-15 顏宗宏", "6-16 林獻崇", "6-17 張淑真", "6-18 許婉瑜", "6-19 劉淑美", "6-20 林俞彣", "6-21 許婷婷"
        ]
    },
    "第 7 組": {
        "leader": "陳定群",
        "members": [
            "7-1 陳定群", "7-2 吳庭妤", "7-3 吳幸蓉", "7-4 林伶雯", "7-5 林美芳", "7-6 林靖芳", "7-7 徐續仁", "7-8 張靜娟", "7-9 蕭淑勻", "7-10 吳璧合", "7-11 陳淑芳", "7-12 黃金輝", "7-13 鄭秀娟", "7-14 鄭宜涵", "7-15 鄭斐文", "7-16 謝萬蒲", "7-17 魏媛真", "7-18 楊美蘭"
        ]
    },
    "第 8 組": {
        "leader": "楊婉君",
        "members": [
            "8-1 楊婉君", "8-2 林宗平", "8-3 吳欣緯", "8-4 陳秋芬", "8-5 王怡誠", "8-6 王麗景", "8-7 章玲珠", "8-8 李瓊華", "8-9 劉燕冰", "8-10 盧菀萱", "8-11 李淑瑛", "8-12 楊素粉", "8-13 劉筱愛", "8-14 盧彥妤", "8-15 明子又", "8-16 黃翠月", "8-17 賴香雪", "8-18 蘇郁合", "8-19 蘇國樑"
        ]
    },
    "第 9 組": {
        "leader": "林如玉",
        "members": [
            "9-1 林如玉", "9-2 林淑娟", "9-3 余雅玲", "9-4 李慧瑛", "9-5 辛天送", "9-6 白秀清", "9-7 莊淑娟", "9-8 游建邦", "9-9 黃惠卿", "9-10 蔡純慧", "9-11 洪雅玲", "9-12 李宛融", "9-13 徐惠美", "9-14 周珈漩", "9-15 許湘惠", "9-16 周林美麗", "9-17 何苑色", "9-18 李俊宏", "9-19 喻榮華", "9-20 柳書玉", "9-21 蘇鴻濱"
        ]
    },
    "第 10 組": {
        "leader": "黃振育",
        "members": [
            "10-1 黃振育", "10-2 胡詠茹", "10-3 李宜晏", "10-4 楊崇誠", "10-5 陳世平", "10-6 張金爗", "10-7 陳麗華", "10-8 廖曉萍", "10-9 賴志雄", "10-10 吳東昱", "10-11 趙如薰", "10-12 叢思瑋", "10-13 李學璁", "10-14 黃品齊", "10-15 朱重信", "10-16 林明珠", "10-17 蔡春美", "10-18 楊昌學", "10-19 謝佩穎", "10-20 葉奕新"
        ]
    },
    "第 11 組": {
        "leader": "林美華",
        "members": [
            "11-1 賴素貞", "11-2 賴彥帛", "11-3 洪于晴", "11-4 何如甘", "11-5 李素茹", "11-6 林珍安", "11-7 吳核豫", "11-8 葉如玲", "11-9 林保秀", "11-10 陳品樺", "11-11 蘇相瑜", "11-12 江盈潔", "11-13 阮芬", "11-14 雙東華", "11-15 林美華", "11-16 蔡菁雯", "11-17 楊金珠", "11-18 蕭淑慧", "11-19 韋淑媛", "11-20 徐丞億"
        ]
    },
    "第 12 組": {
        "leader": "邱品瑄",
        "members": [
            "12-1 邱品瑄", "12-2 賴盈達", "12-3 陳雅惠", "12-4 江佳怡", "12-5 孫詩旻", "12-6 謝秀枝", "12-7 周素仰", "12-8 陳昭華", "12-9 簡達謙", "12-10 潘志滿", "12-11 童小鈴", "12-12 蔡宗涵", "12-13 蔡麗滿", "12-14 王冠蕙", "12-15 洪唯", "12-16 何淑琴", "12-17 顏妙如", "12-18 楊紫璿", "12-19 李淑芬", "12-20 洪金城"
        ]
    },
    "第 13 組": {
        "leader": "高韻筑",
        "members": [
            "13-1 陳重仰", "13-2 高韻筑", "13-3 王上苹", "13-4 陳珈暐", "13-5 林珮圻", "13-6 施瀞雅", "13-7 丁秌全", "13-8 黃蔚菁", "13-9 吳孟霖", "13-10 張慈方", "13-11 鄭鈞懿", "13-12 王慧玟", "13-13 彭美珠", "13-14 黃忠權", "13-15 陳彥伶", "13-16 卓卉綺", "13-17 林于煒", "13-18 張婉玲", "13-19 朱庭葦"
        ]
    },
    "第 14 組": {
        "leader": "蔡秉諺",
        "members": [
            "14-1 蔡秉諺", "14-2 高翊綺", "14-3 秦育曐", "14-4 莊媛婷", "14-5 蔡明真", "14-6 陳桂宸", "14-7 黃麗燕", "14-8 吳承璇", "14-9 葉美鴻", "14-10 張在增", "14-11 顏少萱", "14-12 王馨儀", "14-13 林奕瑄", "14-14 郭芷萱", "14-15 吳靜宜", "14-16 張容", "14-17 蕭米佋", "14-18 李岳璋", "14-19 廖育君", "14-20 謝孟君", "14-21 葉姿佑"
        ]
    },
    "第 15 組": {
        "leader": "高忠偉",
        "members": [
            "15-1 廖以萱", "15-2 高忠偉", "15-3 侯惠芸", "15-4 侯惠倫", "15-5 莊郁琳", "15-6 戴毓書", "15-7 張家瑄", "15-8 邵杰", "15-9 白豐銘", "15-10 黃沛瑋", "15-11 張雁菁", "15-12 蔡佩君", "15-13 鄭伃倢", "15-14 曾品綸", "15-15 張文馨", "15-16 劉思嫺", "15-17 黃姿蓉", "15-18 賴姳璇", "15-19 洪瑩蓁", "15-20 陳星妤"
        ]
    },
    "第 16 組": {
        "leader": "偕魁元",
        "members": [
            "16-1 偕魁元", "16-2 李庭安", "16-3 劉妙娟", "16-4 朱穆鳳", "16-5 沈承宗", "16-6 蔡雅雅", "16-7 凃文鳳", "16-8 莊雅惠", "16-9 劉巧領", "16-10 鄭雅文", "16-11 湯玉琦", "16-12 劉姿吟", "16-13 蕭茲方", "16-14 王芳珠", "16-15 陳冠同", "16-16 林秀菁", "16-17 黃政和", "16-18 徐友信", "16-19 吳莉惠", "16-20 林宜筠"
        ]
    },
    "第 17 組": {
        "leader": "蘇誌盈",
        "members": [
            "17-1 蘇誌盈", "17-2 王素敏", "17-3 蔡石福地", "17-4 李世林", "17-5 黃淑芳", "17-6 胡瑜玲", "17-7 柯孟宜", "17-8 林靜香", "17-9 葉瓊雅", "17-10 張秋美", "17-11 劉宇容", "17-12 黃克經", "17-13 林淑樺", "17-14 黃榮茂", "17-15 邵士誠", "17-16 鄭人豪", "17-17 姚翠華"
        ]
    },
    "第 18 組": {
        "leader": "陳麗華",
        "members": [
            "18-1 陳麗華", "18-2 廖述強", "18-3 黃淑梅", "18-4 林玲芬", "18-5 曹曼容", "18-6 柯錦鑫", "18-7 張湘瑜", "18-8 吳美蓉", "18-9 陳月美", "18-10 陳素葉", "18-11 郭宇家", "18-12 黃佩琪", "18-13 朱龍昌", "18-14 謝桂红", "18-15 阮善如", "18-16 陳丁清", "18-17 楊月娥", "18-18 謝秀錦", "18-19 陳慶旺", "18-20 陳錦鳳"
        ]
    }
}
# Fixed typos / mapping adjustments
groups_data["第 1 組"]["members"][7] = "1-8 郭瓈灧"
groups_data["第 11 組"]["members"][10] = "11-11 蘇相瑜"
groups_data["第 11 組"]["members"][13] = "11-14 許東華"


# Deduplicate members within each group
for g, data in groups_data.items():
    seen = set()
    deduped_members = []
    for m in data["members"]:
        if m not in seen:
            seen.add(m)
            deduped_members.append(m)
    groups_data[g]["members"] = deduped_members

# 週回報期間選項 (1 ~ 40 週)
weeks_list = [
    "W1 ( 2026/08/31 ~ 2026/09/05 )",
    "W2 ( 2026/09/06 ~ 2026/09/12 )",
    "W3 ( 2026/09/13 ~ 2026/09/19 )",
    "W4 ( 2026/09/20 ~ 2026/09/26 )",
    "W5 ( 2026/09/27 ~ 2026/10/03 )",
    "W6 ( 2026/10/04 ~ 2026/10/10 )",
    "W7 ( 2026/10/11 ~ 2026/10/17 )",
    "W8 ( 2026/10/18 ~ 2026/10/24 )",
    "W9 ( 2026/10/25 ~ 2026/10/31 )",
    "W10 ( 2026/11/01 ~ 2026/11/07 )",
    "W11 ( 2026/11/08 ~ 2026/11/14 )",
    "W12 ( 2026/11/15 ~ 2026/11/21 )",
    "W13 ( 2026/11/22 ~ 2026/11/28 )",
    "W14 ( 2026/11/29 ~ 2026/12/05 )",
    "W15 ( 2026/12/06 ~ 2026/12/12 )",
    "W16 ( 2026/12/13 ~ 2026/12/19 )",
    "W17 ( 2026/12/20 ~ 2026/12/26 )",
    "W18 ( 2026/12/27 ~ 2027/01/02 )",
    "W19 ( 2027/01/03 ~ 2027/01/09 )",
    "W20 ( 2027/01/10 ~ 2027/01/16 )",
    "W21 ( 2027/01/17 ~ 2027/01/23 )",
    "W22 ( 2027/01/24 ~ 2027/01/30 )",
    "W23 ( 2027/01/31 ~ 2027/02/06 )",
    "W24 ( 2027/02/07 ~ 2027/02/13 )",
    "W25 ( 2027/02/14 ~ 2027/02/20 )",
    "W26 ( 2027/02/21 ~ 2027/02/27 )",
    "W27 ( 2027/02/28 ~ 2027/03/06 )",
    "W28 ( 2027/03/07 ~ 2027/03/13 )",
    "W29 ( 2027/03/14 ~ 2027/03/20 )",
    "W30 ( 2027/03/21 ~ 2027/03/27 )",
    "W31 ( 2027/03/28 ~ 2027/04/03 )",
    "W32 ( 2027/04/04 ~ 2027/04/10 )",
    "W33 ( 2027/04/11 ~ 2027/04/17 )",
    "W34 ( 2027/04/18 ~ 2027/04/24 )",
    "W35 ( 2027/04/25 ~ 2027/05/01 )",
    "W36 ( 2027/05/02 ~ 2027/05/08 )",
    "W37 ( 2027/05/09 ~ 2027/05/15 )",
    "W38 ( 2027/05/16 ~ 2027/05/22 )",
    "W39 ( 2027/05/23 ~ 2027/05/29 )",
    "W40 ( 2027/05/30 ~ 2027/06/05 )"
]

def get_current_week_index(weeks):
    try:
        today = datetime.datetime.now(pytz.timezone('Asia/Taipei')).date()
        for i, wk_str in enumerate(weeks):
            match = re.search(r'(\d{4}/\d{2}/\d{2})\s*~\s*(\d{4}/\d{2}/\d{2})', wk_str)
            if match:
                s_date = datetime.datetime.strptime(match.group(1), '%Y/%m/%d').date()
                e_date = datetime.datetime.strptime(match.group(2), '%Y/%m/%d').date()
                if s_date <= today <= e_date:
                    return i
        match0 = re.search(r'(\d{4}/\d{2}/\d{2})', weeks[0])
        if match0:
            first_date = datetime.datetime.strptime(match0.group(1), '%Y/%m/%d').date()
            if today < first_date:
                return 0
        return len(weeks) - 1
    except Exception:
        return 0

default_week_idx = get_current_week_index(weeks_list)


# Path to save data
# Path to save data
DB_FILE = "emotional_reports.csv"

# Global Cloud Sync configuration (Reads from Streamlit Secrets automatically!)
gas_api_url = ""
try:
    if "gas_url" in st.secrets:
        gas_api_url = st.secrets["gas_url"]
except Exception:
    pass

# Function to load database
def load_data():
    if os.path.exists(DB_FILE):
        try:
            df = pd.read_csv(DB_FILE)
            if not df.empty and "時間戳記" in df.columns and "個人姓名" in df.columns and "當週回報期間" in df.columns:
                # 依時間戳記排序，確保最後提交的在最後面
                df = df.sort_values(by="時間戳記")
                # 排除重複填報：每個組員在同一週（當週回報期間）只能保留最後一筆（最新）資料，防止統計數據虛胖
                df = df.drop_duplicates(subset=["組別", "當週回報期間", "個人姓名"], keep="last")
            return df
        except Exception:
            pass
    # Return empty template dataframe
    return pd.DataFrame(columns=[
        "時間戳記", "組別", "保溫組長", "個人姓名", "電子郵件", 
        "個人編號", "性別", "當週回報期間", "當週填表張數", 
        "累積填表張數", "本週組內共學", "心得或對法師提問",
        "最常出現的身體反應", "最常出現的情緒類別", "最常用的落地方式"
    ])

# Function to save record (local + cloud try)
def save_record(new_row, gas_url=""):
    df = load_data()
    new_df = pd.DataFrame([new_row])
    df = pd.concat([df, new_df], ignore_index=True)
    df.to_csv(DB_FILE, index=False, encoding='utf-8-sig')
    
    if gas_url and gas_url.strip():
        try:
            response = requests.post(gas_url.strip(), json=new_row, timeout=8)
            if response.status_code == 200:
                return True, "本機備份成功，且順利同步寫入 Google 雲端！"
            else:
                return True, f"本機備份成功，但雲端同步失敗 (狀態碼: {response.status_code})"
        except Exception as e:
            return True, f"本機備份成功，但雲端同步發生錯誤: {str(e)}"
    return True, "填報紀錄已成功儲存於本地伺服器。"

# Define passwords for group leaders (leader01 ~ leader18) and admin (admin2026)
leader_passwords = {f"第 {i} 組": f"leader{i:02d}" for i in range(1, 19)}
admin_password = "admin2026"


import zipfile

def summarize_to_one_sentence(feedback_text):
    if not feedback_text:
        return ""
    text = str(feedback_text).strip()
    if text in ["無", "nan", "None", "無提問", "尚無"]:
        return ""
    
    # Clean newlines and extra spaces
    text = re.sub(r'[\r\n]+', '，', text)
    text = re.sub(r'\s+', ' ', text).strip()
    
    # Split by sentence stops
    sentences = re.split(r'([。！？!])', text)
    
    full_sentences = []
    for i in range(0, len(sentences)-1, 2):
        s = sentences[i].strip()
        p = sentences[i+1]
        if s:
            full_sentences.append(s + p)
    if len(sentences) % 2 == 1 and sentences[-1].strip():
        s = sentences[-1].strip()
        if not (s.endswith('。') or s.endswith('！') or s.endswith('？')):
            s += '。'
        full_sentences.append(s)
        
    if not full_sentences:
        full_sentences = [text + ('。' if not text.endswith('。') else '')]
        
    first_sent = full_sentences[0]
    if len(first_sent) < 15 and len(full_sentences) > 1:
        first_sent = first_sent.rstrip('。！？') + '，' + full_sentences[1]
        
    if len(first_sent) > 85:
        first_sent = first_sent[:80] + "..."
        if not first_sent.endswith("。"):
            first_sent += "。"
            
    return first_sent

def generate_18_groups_zip_reports(week, all_data, groups_data):
    target_wk_label = week.split(' (')[0]
    wk_sub_all = all_data[all_data["當週回報期間"] == week] if not all_data.empty else pd.DataFrame()
    
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for g_name, g_info in groups_data.items():
            ldr = g_info["leader"]
            group_roster = g_info["members"]
            g_num = g_name.replace("第 ", "").replace(" 組", "")
            
            roster_list = []
            for idx, m_full in enumerate(group_roster):
                if " " in m_full:
                    parts = m_full.split(" ")
                    m_id = parts[0]
                    m_name = parts[1]
                else:
                    m_id = f"{g_num}-{idx+1}"
                    m_name = m_full
                roster_list.append({"組員號碼": m_id, "個人姓名": m_name, "編號排序": idx+1})
                
            g_submissions = wk_sub_all[wk_sub_all["組別"] == g_name] if not wk_sub_all.empty else pd.DataFrame()
            
            doc_bio = generate_docx_report(
                group_name=g_name,
                leader=ldr,
                week=week,
                submissions=g_submissions,
                roster_list=roster_list
            )
            
            filename_in_zip = f"{g_name}_情緒觀察每週心得報告_{target_wk_label}.docx"
            zip_file.writestr(filename_in_zip, doc_bio.getvalue())
            
    zip_buffer.seek(0)
    return zip_buffer


def build_group_summary_text(g_name, leader, group_submissions, roster_list):
    """
    Generates a ~200-word synthesized summary paragraph for a group's weekly emotional observation.
    """
    total_m = len(roster_list)
    sub_count = len(group_submissions)
    
    feedbacks = []
    bodies = []
    emotions = []
    groundings = []
    
    if not group_submissions.empty:
        for _, row in group_submissions.iterrows():
            fb = str(row.get("心得或對法師提問", "")).strip()
            bd = str(row.get("最常出現的身體反應", "")).strip()
            em = str(row.get("最常出現的情緒類別", "")).strip()
            gr = str(row.get("最常用的落地方式", "")).strip()
            
            if fb and fb not in ["無", "nan", "None"]:
                feedbacks.append(fb)
            if bd and bd not in ["無", "nan", "None"]:
                bodies.append(bd)
            if em and em not in ["無", "nan", "None"]:
                emotions.append(em)
            if gr and gr not in ["無", "nan", "None"]:
                groundings.append(gr)

    body_str = "、".join(list(set(bodies))[:3]) if bodies else "肌肉緊繃、胸悶或呼吸加速"
    emotion_str = "、".join(list(set(emotions))[:3]) if emotions else "焦慮、嗔怒、不安或浮躁"
    grounding_str = "、".join(list(set(groundings))[:3]) if groundings else "深呼吸、拜佛、觀功念恩與安住當下"
    
    if feedbacks:
        combined_fb = "；".join(feedbacks[:4])
        if len(combined_fb) > 130:
            combined_fb = combined_fb[:125] + "..."
        summary_text = (
            f"本週{g_name}在組長{leader}的溫馨帶領下，共有 {sub_count} 位老師完成情緒觀察週總結填報。"
            f"同修們在日常生活與修持中，最常覺察到的身體反應包含「{body_str}」，相應產生的情緒狀態則以「{emotion_str}」為主；"
            f"面對身心起伏，老師們多能運用「{grounding_str}」等善巧方法自我落地與安撫。在心得與提問方面，同修主要分享了：{combined_fb}。"
            f"整體展現了極佳的向內觀照功力與對法的希求，祈願大家持續在當下練習收攝，累積無量福智資糧。"
        )
    else:
        summary_text = (
            f"本週{g_name}在保溫組長{leader}的悉心護持下，持續推進情緒觀察表的寫作與覺察。"
            f"雖然本週部分老師公務繁忙或尚在整理觀照心得，但組內同修平日多能關注「{body_str}」等身體反應，"
            f"並在面對「{emotion_str}」等心靈起伏時，積極透過「{grounding_str}」來調柔身心、安住當下。"
            f"鼓勵各位老師在忙碌之餘，隨喜自己完成的每一個微小步履，撥空上記記錄，將向內觀察轉化為最真實的修行力量，共同在團隊中增長智慧與慈悲。"
        )
        
    return summary_text

def generate_all_groups_docx_report(week, all_data, groups_data):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft JhengHei'
    font.size = Pt(11)
    
    # 1. Main Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(f"📋 園區情緒觀察營 — 全營 18 組每週總結報告")
    run_title.font.name = 'Microsoft JhengHei'
    run_title.font.size = Pt(18)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(27, 94, 32)
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    today_str = datetime.datetime.now(pytz.timezone('Asia/Taipei')).date().strftime('%Y/%m/%d')
    run_meta = p_meta.add_run(f"回報週別：{week}  |  報表產出日期：{today_str}  |  彙整單位：園區情緒觀察營大組窗口")
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = RGBColor(100, 110, 100)
    
    doc.add_paragraph()
    
    # 2. Executive Summary Table Across All 18 Groups
    h_sum = doc.add_paragraph()
    run_h_sum = h_sum.add_run("📊 全營 18 組填寫率快速評估")
    run_h_sum.font.size = Pt(14)
    run_h_sum.bold = True
    run_h_sum.font.color.rgb = RGBColor(27, 94, 32)
    
    cols_exec = [
        ("組別", Inches(0.65)), 
        ("保溫組長", Inches(1.0)), 
        ("應填人數", Inches(0.6)), 
        ("已填人數", Inches(0.6)), 
        ("未填人數", Inches(0.6)), 
        ("共學人數", Inches(0.65)), 
        ("當週填表張數", Inches(0.85)), 
        ("平均每人填表數", Inches(0.95)), 
        ("填報率", Inches(0.65))
    ]
    table_exec = doc.add_table(rows=1, cols=9)
    table_exec.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_exec.autofit = False
    
    hdr_cells = table_exec.rows[0].cells
    for i, (name, width) in enumerate(cols_exec):
        hdr_cells[i].text = name
        hdr_cells[i].width = width
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.name = 'Microsoft JhengHei'
            r.font.size = Pt(10)
            
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '1B5E20')
        shading.set(qn('w:val'), 'clear')
        tcPr.append(shading)
        
    wk_sub_all = all_data[all_data["當週回報期間"] == week] if not all_data.empty else pd.DataFrame()
    
    total_camp_m = 0
    total_camp_sub = 0
    total_camp_unsub = 0
    total_camp_sheets = 0
    total_camp_co_learn = 0
    
    for g_name, g_info in groups_data.items():
        ldr = g_info["leader"]
        m_list_full = g_info["members"]
        m_list_names = [m.split(" ")[1] if " " in m else m for m in m_list_full]
        
        g_sub = wk_sub_all[wk_sub_all["組別"] == g_name] if not wk_sub_all.empty else pd.DataFrame()
        sub_names = g_sub["個人姓名"].tolist() if not g_sub.empty else []
        
        sub_count = len(set(sub_names) & set(m_list_names))
        unsub_count = len(m_list_full) - sub_count
        rate = (sub_count / len(m_list_full)) * 100 if len(m_list_full) > 0 else 0
        g_sheets = int(g_sub["當週填表張數"].sum()) if not g_sub.empty else 0
        g_co_learn = len(g_sub[g_sub["本週組內共學"] == "是"]) if not g_sub.empty and "本週組內共學" in g_sub.columns else 0
        
        avg_sheets = (g_sheets / len(m_list_full)) if len(m_list_full) > 0 else 0.0
        
        total_camp_m += len(m_list_full)
        total_camp_sub += sub_count
        total_camp_unsub += unsub_count
        total_camp_sheets += g_sheets
        total_camp_co_learn += g_co_learn
        
        row_cells = table_exec.add_row().cells
        r_vals = [g_name, ldr, str(len(m_list_full)), str(sub_count), str(unsub_count), f"{g_co_learn} 人", f"{g_sheets} 張", f"{avg_sheets:.1f} 張/人", f"{rate:.1f}%"]
        
        for i, val in enumerate(r_vals):
            row_cells[i].text = val
            row_cells[i].width = cols_exec[i][1]
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Microsoft JhengHei'
                r.font.size = Pt(9.5)
                
    overall_rate = (total_camp_sub / total_camp_m) * 100 if total_camp_m > 0 else 0
    overall_avg = (total_camp_sheets / total_camp_m) if total_camp_m > 0 else 0.0
    p_camp_met = doc.add_paragraph()
    p_camp_met.paragraph_format.space_before = Pt(10)
    p_camp_met.paragraph_format.space_after = Pt(20)
    r_cmet = p_camp_met.add_run(f"💡 全營整體概況：全營 18 小組應填報共 {total_camp_m} 人，當週實際完成登記 {total_camp_sub} 人（其中參加組內共學共 {total_camp_co_learn} 人），未填報 {total_camp_unsub} 人，全營整體填報率為 {overall_rate:.1f}%。當週全營累積填寫情緒觀察表總計 {total_camp_sheets} 張，全營平均每人填寫 {overall_avg:.1f} 張。")
    r_cmet.font.size = Pt(10.5)
    r_cmet.font.bold = True
    r_cmet.font.color.rgb = RGBColor(46, 125, 50)
    
    # 3. Group Detailed Sections
    h_group_sec = doc.add_paragraph()
    run_h_group_sec = h_group_sec.add_run("💬 十八組各組繳交狀況表格與組員心得條列摘要")
    run_h_group_sec.font.size = Pt(14)
    run_h_group_sec.bold = True
    run_h_group_sec.font.color.rgb = RGBColor(27, 94, 32)
    h_group_sec.paragraph_format.space_after = Pt(12)
    
    for g_name, g_info in groups_data.items():
        ldr = g_info["leader"]
        group_roster = g_info["members"]
        g_num = g_name.replace("第 ", "").replace(" 組", "")
        
        roster_list = []
        for idx, m_full in enumerate(group_roster):
            if " " in m_full:
                parts = m_full.split(" ")
                m_id = parts[0]
                m_name = parts[1]
            else:
                m_id = f"{g_num}-{idx+1}"
                m_name = m_full
            roster_list.append({"組員號碼": m_id, "個人姓名": m_name, "編號排序": idx+1})
            
        g_submissions = wk_sub_all[wk_sub_all["組別"] == g_name] if not wk_sub_all.empty else pd.DataFrame()
        sub_map = {row["個人姓名"]: row for _, row in g_submissions.iterrows()} if not g_submissions.empty else {}
        
        # Group Header
        p_gh = doc.add_paragraph()
        p_gh.paragraph_format.space_before = Pt(14)
        p_gh.paragraph_format.space_after = Pt(4)
        r_gh = p_gh.add_run(f"🔴 {g_name} （保溫組長：{ldr}）")
        r_gh.font.size = Pt(13)
        r_gh.bold = True
        r_gh.font.color.rgb = RGBColor(122, 105, 90)
        
        # Group Status Table (Same as Group Leader Report Table)
        cols_data = [("個人編號", Inches(1.1)), ("組員姓名", Inches(1.1)), ("組內共學", Inches(1.0)), ("填報狀態", Inches(1.0)), ("當週張數", Inches(0.9)), ("累計張數", Inches(0.9))]
        table_g = doc.add_table(rows=1, cols=6)
        table_g.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_g.autofit = False
        
        hdr_cells = table_g.rows[0].cells
        for i, (name, width) in enumerate(cols_data):
            hdr_cells[i].text = name
            hdr_cells[i].width = width
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.name = 'Microsoft JhengHei'
                r.font.size = Pt(10)
                
            tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '1B5E20')
            shading.set(qn('w:val'), 'clear')
            tcPr.append(shading)
            
        g_submitted_count = 0
        g_co_learn_count = 0
        g_total_sheets = 0
        
        for row in roster_list:
            m_name = row["個人姓名"]
            m_id = row["組員號碼"]
            
            row_cells = table_g.add_row().cells
            status = "❌ 未填"
            co_learn_str = "-"
            w_sheets = "-"
            c_sheets = "-"
            
            if m_name in sub_map:
                status = "✅ 已填"
                g_submitted_count += 1
                co_val = str(sub_map[m_name].get("本週組內共學", "否")).strip()
                if co_val == "是":
                    co_learn_str = "✅ 是"
                    g_co_learn_count += 1
                else:
                    co_learn_str = "❌ 否"
                try:
                    w_sheets = str(int(sub_map[m_name]["當週填表張數"]))
                    g_total_sheets += int(sub_map[m_name]["當週填表張數"])
                except Exception:
                    w_sheets = "0"
                try:
                    c_sheets = str(int(sub_map[m_name]["累積填表張數"]))
                except Exception:
                    c_sheets = "0"
                
            row_data = [m_id, m_name, co_learn_str, status, w_sheets, c_sheets]
            
            for i, val in enumerate(row_data):
                row_cells[i].text = val
                row_cells[i].width = cols_data[i][1]
                p = row_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Microsoft JhengHei'
                    r.font.size = Pt(9.5)
                    if status == "❌ 未填" and i == 2:
                        r.font.color.rgb = RGBColor(183, 28, 28)
                    elif status == "✅ 已填" and i == 2:
                        r.font.color.rgb = RGBColor(46, 125, 50)
                        
        p_gmet = doc.add_paragraph()
        p_gmet.paragraph_format.space_before = Pt(6)
        p_gmet.paragraph_format.space_after = Pt(8)
        r_gm = p_gmet.add_run(f"💡 小組小結：應填 {len(roster_list)} 人，已填 {g_submitted_count} 人（組內共學出席 {g_co_learn_count} 人），未填 {len(roster_list) - g_submitted_count} 人。當週小組總填表張數：{g_total_sheets} 張。")
        r_gm.font.size = Pt(9.5)
        r_gm.font.bold = True
        r_gm.font.color.rgb = RGBColor(100, 100, 100)
        
        # Group Member Thoughts Itemized Summary (Direct 1-sentence summaries, skipping unwritten members)
        p_sum_lbl = doc.add_paragraph()
        p_sum_lbl.paragraph_format.space_before = Pt(6)
        p_sum_lbl.paragraph_format.space_after = Pt(4)
        r_s_lbl = p_sum_lbl.add_run(f"📝 {g_name} 本週組員心得條列摘要：")
        r_s_lbl.bold = True
        r_s_lbl.font.size = Pt(10.5)
        r_s_lbl.font.color.rgb = RGBColor(27, 94, 32)
        
        has_written = False
        for row in roster_list:
            m_name = row["個人姓名"]
            m_id = row["組員號碼"]
            if m_name in sub_map:
                raw_fb = sub_map[m_name].get("心得或對法師提問", "")
                one_sentence = summarize_to_one_sentence(raw_fb)
                if one_sentence:
                    has_written = True
                    p_item = doc.add_paragraph()
                    p_item.paragraph_format.left_indent = Inches(0.2)
                    p_item.paragraph_format.space_after = Pt(3)
                    
                    r_id_name = p_item.add_run(f"• 【{m_id} {m_name}】：")
                    r_id_name.bold = True
                    r_id_name.font.size = Pt(10)
                    r_id_name.font.color.rgb = RGBColor(78, 52, 46)
                    
                    r_txt = p_item.add_run(one_sentence)
                    r_txt.font.size = Pt(10)
                    r_txt.font.color.rgb = RGBColor(40, 40, 40)
                    
        if not has_written:
            p_none = doc.add_paragraph()
            p_none.paragraph_format.left_indent = Inches(0.2)
            p_none.paragraph_format.space_after = Pt(10)
            r_none = p_none.add_run("（本週貴組組員暫無填寫觀照心得或提問）")
            r_none.font.size = Pt(9.5)
            r_none.italic = True
            r_none.font.color.rgb = RGBColor(140, 140, 140)
        
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_after = Pt(6)
        r_div = p_div.add_run("―" * 55)
        r_div.font.color.rgb = RGBColor(220, 220, 220)
        r_div.font.size = Pt(8)
        
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# Helper to generate beautiful Word Document (.docx) Weekly Thoughts Report
def generate_docx_report(group_name, leader, week, submissions, roster_list):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft JhengHei'
    font.size = Pt(11)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(f"📋 情緒觀察營 — {group_name} 每週心得報告")
    run_title.font.name = 'Microsoft JhengHei'
    run_title.font.size = Pt(18)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(27, 94, 32)
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run(f"回報週別：{week}  |  保溫組長：{leader}  |  報表產出日期：{datetime.datetime.now(pytz.timezone('Asia/Taipei')).date().strftime('%Y/%m/%d')}")
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = RGBColor(100, 110, 100)
    
    doc.add_paragraph()
    
    h_sum = doc.add_paragraph()
    run_h_sum = h_sum.add_run("📊 當週填表數據摘要表")
    run_h_sum.font.size = Pt(14)
    run_h_sum.bold = True
    run_h_sum.font.color.rgb = RGBColor(27, 94, 32)
    
    cols_data = [("個人編號", Inches(1.1)), ("組員姓名", Inches(1.1)), ("組內共學", Inches(1.0)), ("填報狀態", Inches(1.0)), ("當週張數", Inches(0.9)), ("累計張數", Inches(0.9))]
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    hdr_cells = table.rows[0].cells
    for i, (name, width) in enumerate(cols_data):
        hdr_cells[i].text = name
        hdr_cells[i].width = width
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.name = 'Microsoft JhengHei'
            r.font.size = Pt(10.5)
            
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '1B5E20')
        shading.set(qn('w:val'), 'clear')
        tcPr.append(shading)
        
    sub_map = {row["個人姓名"]: row for _, row in submissions.iterrows()}
    
    total_sheets = 0
    submitted_count = 0
    co_learn_count = 0
    
    for row in roster_list:
        m_name = row["個人姓名"]
        m_id = row["組員號碼"]
        
        row_cells = table.add_row().cells
        status = "❌ 未填"
        co_learn_str = "-"
        w_sheets = "-"
        c_sheets = "-"
        
        if m_name in sub_map:
            status = "✅ 已填"
            submitted_count += 1
            co_val = str(sub_map[m_name].get("本週組內共學", "否")).strip()
            if co_val == "是":
                co_learn_str = "✅ 是"
                co_learn_count += 1
            else:
                co_learn_str = "❌ 否"
            w_sheets = str(int(sub_map[m_name]["當週填表張數"]))
            c_sheets = str(int(sub_map[m_name]["累積填表張數"]))
            total_sheets += int(sub_map[m_name]["當週填表張數"])
            
        row_data = [m_id, m_name, co_learn_str, status, w_sheets, c_sheets]
        
        for i, val in enumerate(row_data):
            row_cells[i].text = val
            row_cells[i].width = cols_data[i][1]
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Microsoft JhengHei'
                r.font.size = Pt(10)
                if status == "❌ 未填" and i == 2:
                    r.font.color.rgb = RGBColor(183, 28, 28)
                elif status == "✅ 已填" and i == 2:
                    r.font.color.rgb = RGBColor(46, 125, 50)
                    
    p_metrics = doc.add_paragraph()
    p_metrics.paragraph_format.space_before = Pt(12)
    p_metrics.paragraph_format.space_after = Pt(24)
    r_met = p_metrics.add_run(f"💡 本週小組統計：應填報 {len(roster_list)} 人，已填報 {submitted_count} 人（組內共學出席 {co_learn_count} 人），未填報 {len(roster_list) - submitted_count} 人。當週填表總張數：{total_sheets} 張。")
    r_met.font.size = Pt(10)
    r_met.font.bold = True
    r_met.font.color.rgb = RGBColor(100, 100, 100)
    
    h_detail = doc.add_paragraph()
    run_h_detail = h_detail.add_run("💬 組員心得與對法師提問明細")
    run_h_detail.font.size = Pt(14)
    run_h_detail.bold = True
    run_h_detail.font.color.rgb = RGBColor(27, 94, 32)
    h_detail.paragraph_format.space_after = Pt(12)
    
    for row in roster_list:
        m_name = row["個人姓名"]
        m_id = row["組員號碼"]
        
        p_mem = doc.add_paragraph()
        p_mem.paragraph_format.space_before = Pt(8)
        p_mem.paragraph_format.space_after = Pt(2)
        r_mem = p_mem.add_run(f"【{m_id}】 {m_name}")
        r_mem.bold = True
        r_mem.font.size = Pt(11.5)
        r_mem.font.color.rgb = RGBColor(78, 52, 46)
        
        if m_name in sub_map:
            sub = sub_map[m_name]
            sheets = sub["當週填表張數"]
            cum_sheets = sub["累積填表張數"]
            feedback = sub["心得或對法師提問"]
            
            p_sheets = doc.add_paragraph()
            p_sheets.paragraph_format.left_indent = Inches(0.25)
            p_sheets.paragraph_format.space_after = Pt(2)
            co_val = str(sub.get("本週組內共學", "否")).strip()
            co_txt = "是" if co_val == "是" else "否"
            r_sh = p_sheets.add_run(f"📊 本週填表：{int(sheets)} 張 | 累計填表：{int(cum_sheets)} 張 | 組內共學：{co_txt}")
            r_sh.font.size = Pt(9.5)
            r_sh.italic = True
            r_sh.font.color.rgb = RGBColor(120, 120, 120)
            
            # 💡 v25 新增三欄數據呈現 (若組員有填寫)
            extra_lines = []
            if "最常出現的身體反應" in sub and str(sub["最常出現的身體反應"]).strip() not in ["", "nan", "None", "無"]:
                extra_lines.append(f"肢體反應: {sub['最常出現的身體反應']}")
            if "最常出現的情緒類別" in sub and str(sub["最常出現的情緒類別"]).strip() not in ["", "nan", "None", "無"]:
                extra_lines.append(f"情緒類別: {sub['最常出現的情緒類別']}")
            if "最常用的落地方式" in sub and str(sub["最常用的落地方式"]).strip() not in ["", "nan", "None", "無"]:
                extra_lines.append(f"落地方式: {sub['最常用的落地方式']}")
                
            if extra_lines:
                p_extra = doc.add_paragraph()
                p_extra.paragraph_format.left_indent = Inches(0.25)
                p_extra.paragraph_format.space_after = Pt(2)
                r_ex = p_extra.add_run(" | ".join(extra_lines))
                r_ex.font.size = Pt(9.5)
                r_ex.font.color.rgb = RGBColor(100, 110, 100)
            
            p_feed = doc.add_paragraph()
            p_feed.paragraph_format.left_indent = Inches(0.25)
            p_feed.paragraph_format.space_after = Pt(12)
            
            if str(feedback).strip() in ["", "無", "nan", "None"]:
                r_feed = p_feed.add_run("（組員本週無填寫心得或提問）")
                r_feed.font.color.rgb = RGBColor(150, 150, 150)
                r_feed.italic = True
            else:
                r_feed = p_feed.add_run(str(feedback).strip())
                r_feed.font.color.rgb = RGBColor(33, 33, 33)
                
        else:
            p_feed = doc.add_paragraph()
            p_feed.paragraph_format.left_indent = Inches(0.25)
            p_feed.paragraph_format.space_after = Pt(12)
            r_feed = p_feed.add_run("❌ 本週尚未提交週總結問卷。")
            r_feed.font.color.rgb = RGBColor(183, 28, 28)
            r_feed.italic = True
            
        p_div = doc.add_paragraph()
        p_div.paragraph_format.space_before = Pt(4)
        p_div.paragraph_format.space_after = Pt(4)
        r_div = p_div.add_run("―" * 50)
        r_div.font.color.rgb = RGBColor(230, 230, 230)
        r_div.font.size = Pt(8)
        
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# Helper to send a beautifully formatted confirmation email copy to students
def send_email_copy(receiver_email, row_data):
    # Check if SMTP details are in Secrets
    smtp_user = ""
    smtp_password = ""
    smtp_server = "smtp.gmail.com"
    smtp_port = 465
    
    try:
        if "smtp_user" in st.secrets:
            smtp_user = st.secrets["smtp_user"]
        if "smtp_password" in st.secrets:
            smtp_password = st.secrets["smtp_password"]
        if "smtp_server" in st.secrets:
            smtp_server = st.secrets["smtp_server"]
        if "smtp_port" in st.secrets:
            smtp_port = int(st.secrets["smtp_port"])
    except Exception:
        return False, "無法讀取 Streamlit Secrets 中的 SMTP 設定。"
        
    if not smtp_user or not smtp_password:
        return False, "大組總負責人尚未在 Streamlit Secrets 中設定 'smtp_user' 與 'smtp_password'，因此暫時無法發送電子郵件副本。但您的填報資料已妥善儲存！"
        
    try:
        # Construct beautifully formatted HTML email
        subject = f"【情緒觀察營】填報確認副本 — {row_data['個人姓名']}"
        
        html_content = f"""
        <html>
        <body style="font-family: 'Microsoft JhengHei', sans-serif; color: #2E4436; background-color: #F6FAF4; padding: 20px;">
            <div style="max-width: 600px; margin: 0 auto; background-color: #FFFFFF; padding: 30px; border-radius: 12px; border: 1px solid #D5E5DE; box-shadow: 2px 4px 15px rgba(143, 168, 147, 0.1);">
                <h2 style="color: #3C6E47; border-bottom: 2px solid #8FA893; padding-bottom: 10px; text-align: center;">🌸 園區情緒觀察 填報確認副本</h2>
                <p>親愛的 <b>{row_data['個人姓名']}</b> 吉祥：</p>
                <p>感謝您撥冗填報情緒觀察當週總結！以下是您於 <b>{row_data['時間戳記']}</b> 提交的填報內容，請您留存備查：</p>
                
                <table style="width: 100%; border-collapse: collapse; margin: 20px 0;">
                    <tr style="background-color: #EBF3EC;">
                        <th style="padding: 10px; border: 1px solid #D5E5DE; text-align: left; width: 35%; color: #1B5E20;">項目</th>
                        <th style="padding: 10px; border: 1px solid #D5E5DE; text-align: left; color: #2E4436;">您的回報內容</th>
                    </tr>
                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">報名組別</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE;">{row_data['組別']}</td>
                    </tr>

                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">個人編號</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE;">{row_data['個人編號']}</td>
                    </tr>
                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">當週回報期間</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE;">{row_data['當週回報期間']}</td>
                    </tr>
                    <tr style="background-color: #FAF5EF;">
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #1B5E20;">本週組內共學</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #1B5E20;">{row_data.get('本週組內共學', '否')}</td>
                    </tr>
                    <tr style="background-color: #FAF5EF;">
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #1B5E20;">題目一：當週填表張數</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #1B5E20;">{row_data['當週填表張數']} 張</td>
                    </tr>
                    <tr style="background-color: #FAF5EF;">
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #1B5E20;">題目二：累積填表張數</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #1B5E20;">{row_data['累積填表張數']} 張 (自 2026/8/17 起算)</td>
                    </tr>
                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">2.本週最常出現的身體反應</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE;">{row_data.get('最常出現的身體反應', '無')}</td>
                    </tr>
                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">3.最常出現的情緒類別</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE;">{row_data.get('最常出現的情緒類別', '無')}</td>
                    </tr>
                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">4.最常用的落地方式</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE;">{row_data.get('最常用的落地方式', '無')}</td>
                    </tr>
                    <tr>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; font-weight: bold; color: #7A695A;">5.心得、困難或對法師提問</td>
                        <td style="padding: 10px; border: 1px solid #D5E5DE; white-space: pre-wrap;">{row_data['心得或對法師提問']}</td>
                    </tr>
                </table>
                

                

            </div>
        </body>
        </html>
        """
        
        msg = MIMEText(html_content, "html", "utf-8")
        msg["Subject"] = Header(subject, "utf-8")
        msg["From"] = Header(f"情緒觀察營大組窗口 <{smtp_user}>", "utf-8")
        msg["To"] = Header(receiver_email, "utf-8")
        
        # Connect to server
        if smtp_port == 465:
            server = smtplib.SMTP_SSL(smtp_server, smtp_port, timeout=8)
        else:
            server = smtplib.SMTP(smtp_server, smtp_port, timeout=8)
            server.starttls()
            
        server.login(smtp_user, smtp_password)
        server.sendmail(smtp_user, [receiver_email], msg.as_string())
        server.quit()
        return True, "且已自動將填報確認副本信寄送至您的信箱！"
    except Exception as e:
        return False, f"填報已成功儲存，但發送確認郵件時發生錯誤: {str(e)}"



# Initialize session state variables
if "is_authenticated" not in st.session_state:
    st.session_state["is_authenticated"] = False
if "role" not in st.session_state:
    st.session_state["role"] = "📝 一般學員填寫問卷"
if "logged_in_group" not in st.session_state:
    st.session_state["logged_in_group"] = None
if "submitted_successfully" not in st.session_state:
    st.session_state["submitted_successfully"] = False

# Sidebar Configuration (only visible when logged in)
if st.session_state["is_authenticated"]:
    with st.sidebar:
        st.image("https://img.icons8.com/clouds/200/database.png", width=80)
        st.markdown(f"### 👤 目前登入身分\n**{st.session_state['role']}**")
        if st.session_state["logged_in_group"]:
            st.markdown(f"**管理組別：{st.session_state['logged_in_group']}**")
            

        
        st.write("---")
        # Logout button
        if st.button("🔒 登出管理系統", type="primary", key="sidebar_logout_btn"):
            st.session_state["is_authenticated"] = False
            st.session_state["role"] = "📝 一般學員填寫問卷"
            st.session_state["logged_in_group"] = None
            st.session_state["submitted_successfully"] = False
            # clear query params
            st.query_params.clear()
            st.rerun()
else:
    # If not authenticated, hide sidebar completely!
    st.markdown("""
    <style>
        [data-testid="collapsedControl"] {
            display: none !important;
        }
        section[data-testid="stSidebar"] {
            display: none !important;
        }
    
    /* 強制所有輸入框、文字區域、下拉選單、下拉清單外框使用白/淺色背景與深色文字，確保極高對比度且在手機或暗黑模式下清晰可讀 */
    input, textarea, select, 
    div[data-baseweb="select"] > div, 
    div[data-baseweb="textarea"] > div, 
    div[data-baseweb="input"] > div,
    [role="listbox"], [role="option"], 
    [data-baseweb="popover"] div, 
    [data-baseweb="popover"] ul, 
    [data-baseweb="popover"] li {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-color: #D5E5DE !important;
    }
    /* 當焦點在輸入框或文字區時的文字顏色也必須保持深森林綠 */
    input:focus, textarea:focus {
        color: #2E4436 !important;
        background-color: #FFFFFF !important;
    }

    /* 確保側邊欄（Sidebar）在任何模式下都是溫暖的水彩粘土白，而非黑/深灰色 */
    section[data-testid="stSidebar"], [data-testid="stSidebarUserContent"], div[data-testid="stSidebarCollapseButton"] {
        background-color: #FAF5EF !important;
        background-image: linear-gradient(180deg, #F3F8F2 0%, #FAF5EF 100%) !important;
        color: #2E4436 !important;
    }
    section[data-testid="stSidebar"] p, 
    section[data-testid="stSidebar"] span, 
    section[data-testid="stSidebar"] label, 
    section[data-testid="stSidebar"] h1, 
    section[data-testid="stSidebar"] h2, 
    section[data-testid="stSidebar"] h3 {
        color: #2E4436 !important;
    }

    /* 確保表格 Dataframe 在任何模式下皆為白色背景與深色文字，全置中對齊防看錯行 */
    div[data-testid="stTable"] table, div[data-testid="stDataFrame"] table {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-collapse: separate !important;
        border-spacing: 0 !important;
        width: 100% !important;
    }
    div[data-testid="stTable"] th, div[data-testid="stDataFrame"] th {
        background-color: #EBF3EC !important;
        color: #1B5E20 !important;
        text-align: center !important; /* 標頭置中對齊 */
        vertical-align: middle !important;
    }
    div[data-testid="stTable"] td, div[data-testid="stDataFrame"] td {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-top: 1px solid #EAE0D5 !important; /* 補上頂底邊框線防止 separate 模式下邊框不見 */
        border-bottom: 1px solid #EAE0D5 !important;
        text-align: center !important; /* 數據置中對齊，防手機上看錯行 */
        vertical-align: middle !important;
    }

    /* 💡 v38 調整：全螢幕扁平外框，移除 max-height 與捲動限制，使其變回一般大方表格 */
    div[data-testid="stTable"] {
        overflow-x: auto !important;
        position: relative !important;
        border: 1px solid #EAE0D5 !important;
        border-radius: 12px !important;
        box-shadow: inset 0 1px 3px rgba(0,0,0,0.02) !important;
        background-color: #FFFFFF !important;
        -webkit-overflow-scrolling: touch !important;
    }

    /* v38 新增: 系統全局數據篩選器專屬的「視窗滾動容器」與「頂端列凍結（Sticky Header）」 */
    .scroll-table-container {
        max-height: 450px !important;
        overflow-y: auto !important;
        overflow-x: auto !important;
        position: relative !important;
        border: 1px solid #EAE0D5 !important;
        border-radius: 12px !important;
        box-shadow: inset 0 1px 3px rgba(0,0,0,0.02) !important;
        background-color: #FFFFFF !important;
        -webkit-overflow-scrolling: touch !important; /* 確保 iOS 手機端滑動流暢 */
        margin-top: 10px;
        margin-bottom: 20px;
    }
    .scroll-table-container table {
        width: 100% !important;
        border-collapse: separate !important;
        border-spacing: 0 !important;
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        font-size: 0.95rem;
    }
    /* 保持表頭固定在頂部 */
    .scroll-table-container thead,
    .scroll-table-container thead tr,
    .scroll-table-container th {
        position: sticky !important;
        top: 0 !important;
        z-index: 99 !important; /* 確保在滾動時絕對壓在下方資料之上 */
    }
    .scroll-table-container th {
        background-color: #EBF3EC !important;
        color: #1B5E20 !important;
        text-align: center !important;
        vertical-align: middle !important;
        padding: 10px 15px !important;
        font-weight: bold !important;
        box-shadow: 0 1px 0 #D5E5DE, 0 2px 4px rgba(0,0,0,0.04) !important; /* 加上下分割陰影 */
        border: none !important;
    }
    .scroll-table-container td {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
        border-top: 1px solid #EAE0D5 !important;
        border-bottom: 1px solid #EAE0D5 !important;
        text-align: center !important;
        vertical-align: middle !important;
        padding: 10px 15px !important;
    }

    /* 確保 Metric 數據指標卡片也有高對比度深色文字與白/淺色底色，在任何情況下都不會呈現黑底 */
    div[data-testid="stMetricValue"] > div {
        color: #1B5E20 !important; /* 清新綠數字 */
    }
    div[data-testid="stMetricLabel"] > div {
        color: #7A695A !important; /* 泥土腳標題 */
    }
    div[data-testid="stMetric"] {
        background-color: #FAF5EF !important;
        border-radius: 12px !important;
        padding: 12px !important;
        border: 1px solid #EAE0D5 !important;
        box-shadow: 1px 3px 10px rgba(122, 105, 90, 0.08) !important;
    }


    /* v22 新增：強制程式碼區塊 (st.code) 呈現極緻高對比度淺底深綠字，杜絕黑底 */
    div[data-testid="stCodeBlock"], 
    div[data-testid="stCodeBlock"] pre, 
    div[data-testid="stCodeBlock"] code,
    div[data-testid="stCodeBlock"] span,
    code, pre {
        background-color: #FAF5EF !important;
        color: #2E4436 !important;
        border-radius: 8px !important;
        border: 1px solid #EAE0D5 !important;
    }

    /* v22 新增：強制 Multiselect 選項標籤 (Tag/Chip) 在暗黑模式下為淺底深綠字，杜絕黑底 */
    div[data-baseweb="tag"] {
        background-color: #EBF3EC !important;
        color: #2E4436 !important;
        border: 1px solid #D5E5DE !important;
    }
    div[data-baseweb="tag"] span {
        color: #2E4436 !important;
    }

    /* v22 新增：強制所有輸入、文字區域、下拉選單容器保持高對比度純白底色，在任何情況下都不會呈現黑底 */
    div[data-testid="stTextArea"] textarea, 
    div[data-testid="stTextInput"] input, 
    div[data-testid="stNumberInput"] input, 
    div[data-testid="stSelectbox"] select,
    div[data-testid="stTextArea"] div, 
    div[data-testid="stTextInput"] div, 
    div[data-testid="stNumberInput"] div, 
    div[data-testid="stSelectbox"] div,
    div[data-baseweb="input"], 
    div[data-baseweb="textarea"],
    div[data-baseweb="select"] {
        background-color: #FFFFFF !important;
        color: #2E4436 !important;
    }

</style>
    """, unsafe_allow_html=True)


# ----------------- Check Query Params for Admin Login -----------------
query_role = st.query_params.get("role", "")
query_manage = st.query_params.get("manage", "")

show_admin_login_screen = False
if (query_role in ["manage", "leader", "admin"] or query_manage == "true") and not st.session_state["is_authenticated"]:
    show_admin_login_screen = True

# Get current state variables for standard compatibility
is_authenticated = st.session_state["is_authenticated"]
role = st.session_state["role"]
logged_in_group = st.session_state["logged_in_group"]

# ----------------- Router Logic -----------------
if show_admin_login_screen:
    # 1. Dedicated separate Admin/Leader login view
    st.markdown("<div style='text-align: center; margin-top: 30px;'><img src='https://img.icons8.com/clouds/200/lock.png' width='120'></div>", unsafe_allow_html=True)
    st.subheader("🔐 情緒觀察營 — 管理者專屬登入頁面")
    st.info("💡 請輸入您的組長密碼（例如 leader10）或總負責人密碼（例如 admin2026）解鎖資料閱讀。")
    
    login_pwd = st.text_input("請輸入管理密碼：", type="password", key="dedicated_pwd_input")
    
    col_btn1, col_btn2 = st.columns([1, 4])
    with col_btn1:
        login_click = st.button("確認登入 🔓", type="primary", key="dedicated_login_btn")
    with col_btn2:
        back_to_form = st.button("返回學員填報頁面 📝", key="dedicated_back_btn")
        if back_to_form:
            st.query_params.clear()
            st.rerun()
            
    if login_click and login_pwd:
        if login_pwd == admin_password:
            st.session_state["is_authenticated"] = True
            st.session_state["role"] = "📊 總窗口後台管理"
            st.session_state["logged_in_group"] = None
            st.success("🔓 總窗口解鎖成功！您擁有全局最高權限。")
            st.query_params.clear()
            st.rerun()
        else:
            matched_group = None
            for g, p in leader_passwords.items():
                if login_pwd == p:
                    matched_group = g
                    break
            if matched_group:
                st.session_state["is_authenticated"] = True
                st.session_state["role"] = "👥 保溫組長專區"
                st.session_state["logged_in_group"] = matched_group
                st.success(f"🔓 登入成功！已解鎖【{matched_group}】權限。")
                st.query_params.clear()
                st.rerun()
            else:
                st.error("密碼錯誤")

elif not is_authenticated:
    # 2. Student Questionnaire View (Ultra-clean, No st.form, Zero mobile confirmation alerts)
    st.info("我有一個身體，我有一個人身。    注視當下，就是前進。")
    
    st.subheader("組別與姓名")
    col1, col2 = st.columns(2)
    with col1:
        group_selected = st.selectbox("請選擇您的組別 *", list(groups_data.keys()), index=9) # Default to Group 10
    
    with col2:
        names_list = groups_data[group_selected]["members"]
        name_selected = st.selectbox("請選擇您的姓名 *", names_list)
        
    leader = groups_data[group_selected]["leader"]
    
    group_num = group_selected.replace("第 ", "").replace(" 組", "")
    if name_selected and " " in name_selected:
        parts = name_selected.split(" ")
        auto_member_id = parts[0]
        actual_name = parts[1]
    else:
        auto_member_id = f"{group_num}-X"
        actual_name = name_selected
        
    st.markdown(f'<div class="group-box">ℹ️ 您選取的是 <b>{group_selected}</b>，請確認已選擇編號為：<b>{auto_member_id}</b></div>', unsafe_allow_html=True)
    
    if st.session_state.get("submitted_successfully", False):
        st.markdown(f"""
        <div class="success-box">
            <h3 style="margin-top: 0; color: #0F5132;">🎉 提交成功！</h3>
            <p>感謝 <b>{st.session_state.get('last_submitted_name', '')}</b> 的回報！您的填表紀錄已妥善儲存。</p>
            <p>💡 {st.session_state.get('last_submitted_msg', '')}</p>
            <ul>
                <li><b>個人編號：</b>{st.session_state.get('last_submitted_id', '')}</li>
                <li><b>回報期間：</b>{st.session_state.get('last_submitted_period', '')}</li>
                <li><b>本週組內共學：</b>{st.session_state.get('last_submitted_co_learn', '否')}</li>
                <li><b>當週張數：</b>{st.session_state.get('last_submitted_weekly', 0)} 張</li>
                <li><b>累計張數：</b>{st.session_state.get('last_submitted_cum', 0)} 張</li>
                <li><b>2.本週最常出現的身體反應：</b>{st.session_state.get('last_submitted_body_reactions', '無')}</li>
                <li><b>3.最常出現的情緒類別：</b>{st.session_state.get('last_submitted_emotion_types', '無')}</li>
                <li><b>4.最常用的落地方式：</b>{st.session_state.get('last_submitted_grounding_methods', '無')}</li>
            </ul>
        </div>
        """, unsafe_allow_html=True)
        
        if st.button("再填一筆 📝", type="primary", key="reset_submission_btn"):
            st.session_state["submitted_successfully"] = False
            st.rerun()
    else:
        st.subheader("回報週次")
        period = st.selectbox("請確認當週回報期間 *", weeks_list, index=default_week_idx)
        
        st.write("---")
        st.markdown("#### 👥 本週組內共學")
        co_learn = st.checkbox("我有參加本週組內共學", value=False, key="student_co_learn")
        
        st.write("---")
        st.markdown("#### 1.填表張數")
        
        weekly_sheets = st.number_input("您本週寫了幾張情緒觀察表？ *", min_value=0, max_value=1000, value=0, step=1, key="student_weekly_sheets")
        cum_sheets = st.number_input("填表總數：自 2026/8/17 起累計填寫了幾張情緒觀察表（含本週） *", min_value=0, max_value=9999, value=0, step=1, key="student_cum_sheets")
        
        st.write("---")
        st.markdown("#### 2.本週最常出現的身體反應")
        body_reactions = st.text_input(
            "2.本週最常出現的身體反應",
            placeholder="請簡短填寫您最常出現的身體反應（例如：肌肉緊繃、心跳加快、胸悶、無等）...",
            key="student_body_reactions",
            label_visibility="collapsed"
        )
        
        st.markdown("#### 3.最常出現的情緒類別")
        emotion_types = st.text_input(
            "3.最常出現的情緒類別",
            placeholder="請簡短填寫您最常出現的情緒類別（例如：嗔、傲慢、焦慮、平靜等）...",
            key="student_emotion_types",
            label_visibility="collapsed"
        )
        
        st.markdown("#### 4.最常用的落地方式")
        grounding_methods = st.text_input(
            "4.最常用的落地方式",
            placeholder="請簡短填寫您最常用的落地方式（例如：深呼吸、拜佛、觀功念恩等）...",
            key="student_grounding_methods",
            label_visibility="collapsed"
        )
        
        st.write("---")
        st.markdown("#### 5.心得、困難或對法師提問")
        feedback = st.text_area(
            "5.心得、困難或對法師提問", 
            placeholder="請輸入您當週的觀察心得，或是想對法師提出的疑問...", 
            key="student_feedback",
            label_visibility="collapsed"
        )
        
        st.write("---")
        st.markdown('<div class="reminder-template" style="font-family: inherit; font-size: 1rem; line-height: 1.5;">💡 <b>溫馨提醒：</b>每週六晚上11:00為報截止，可養成習慣總結。麻煩組長在週日中午產生組內報告回傳給帶組法師。</div>', unsafe_allow_html=True)
        st.write("---")
        student_email_val = st.text_input("📧 您的電子郵件 (選填，填寫後送出可自動收到填報確認信副本)", placeholder="username@gmail.com", key="student_email_input")
        
        submit_btn = st.button("提交問卷 📨", type="primary", key="student_submit_btn")
        
        if submit_btn:
            email = student_email_val.strip() if student_email_val.strip() else f"{actual_name}@example.com"
            GENDER_MAP = {
    "丁秌全": "男",
    "何如甘": "女",
    "何淑琴": "女",
    "何苑色": "女",
    "何進德": "男",
    "余雅玲": "女",
    "侯得正": "男",
    "侯惠倫": "女",
    "侯惠芸": "女",
    "偕魁元": "男",
    "凃文鳳": "女",
    "劉以婕": "女",
    "劉俊麟": "男",
    "劉妙娟": "女",
    "劉姿吟": "女",
    "劉宇容": "女",
    "劉巧領": "女",
    "劉德政": "男",
    "劉思嫺": "女",
    "劉洹岑": "女",
    "劉淑美": "女",
    "劉潓雪": "女",
    "劉燕冰": "女",
    "劉玉如": "女",
    "劉筱愛": "女",
    "劉馨檀": "女",
    "卓卉綺": "女",
    "叢思瑋": "女",
    "吳孟霖": "女",
    "吳幸蓉": "女",
    "吳庭妤": "女",
    "吳念真": "女",
    "吳承璇": "女",
    "吳振仁": "男",
    "吳東昱": "男",
    "吳核豫": "女",
    "吳欣緯": "女",
    "吳煥崇": "男",
    "吳璧合": "女",
    "吳美毅": "女",
    "吳美蓉": "女",
    "吳翊菱": "女",
    "吳莉惠": "女",
    "吳靜宜": "女",
    "周文祥": "男",
    "周林美麗": "女",
    "周珈漩": "女",
    "周秀芬": "女",
    "周素仰": "女",
    "喻榮華": "男",
    "嚴惠英": "女",
    "姚翠華": "女",
    "孫詩旻": "女",
    "尤躍勳": "男",
    "廖以萱": "女",
    "廖曉萍": "女",
    "廖育君": "女",
    "廖袖婷": "女",
    "廖述強": "男",
    "廖鶴翔": "男",
    "張又方": "女",
    "張在增": "男",
    "張婉玲": "女",
    "張家瑄": "女",
    "張容": "女",
    "張志豪": "男",
    "張慈方": "女",
    "張文馨": "女",
    "張淑真": "女",
    "張湘瑜": "女",
    "張玉芳": "女",
    "張秋美": "女",
    "張筱雯": "女",
    "張金爗": "男",
    "張開明": "男",
    "張雁菁": "女",
    "張靜娟": "女",
    "張顥嚴": "男",
    "張麗卿": "女",
    "彭美珠": "女",
    "彭聖森": "男",
    "徐丞億": "男",
    "徐友信": "男",
    "徐惠美": "女",
    "徐立昇": "男",
    "徐續仁": "女",
    "戈德": "男",
    "戴世昌": "男",
    "戴毓書": "女",
    "施瀞雅": "女",
    "明子又": "女",
    "曹曼容": "女",
    "曾品綸": "女",
    "曾旭正": "男",
    "曾美芳": "女",
    "曾麗珍": "女",
    "朱庭葦": "女",
    "朱穆鳳": "女",
    "朱重信": "男",
    "朱龍昌": "男",
    "李世林": "男",
    "李俊宏": "男",
    "李俊男": "男",
    "李學璁": "女",
    "李宛融": "女",
    "李宜晏": "女",
    "李岳璋": "男",
    "李平裕": "男",
    "李庭安": "女",
    "李慧瑛": "女",
    "李淑瑛": "女",
    "李淑芬": "女",
    "李濱如": "女",
    "李瓊華": "女",
    "李素茹": "女",
    "李麗卿": "女",
    "李麗斐": "女",
    "林于煒": "女",
    "林伶雯": "女",
    "林佳慧": "女",
    "林保秀": "女",
    "林俞彣": "女",
    "林奕瑄": "女",
    "林如玉": "男",
    "林宗平": "男",
    "林宜筠": "女",
    "林家敬": "男",
    "林德慶": "男",
    "林文華": "男",
    "林明珠": "女",
    "林淑娟": "女",
    "林淑寧": "女",
    "林淑樺": "女",
    "林献崇": "男",
    "林玲芬": "女",
    "林珍安": "女",
    "林珮圻": "女",
    "林秀婷": "女",
    "林秀菁": "女",
    "林美芳": "女",
    "林美華": "女",
    "林翔均": "男",
    "林耘庄": "女",
    "林靖芳": "女",
    "林靜香": "女",
    "林麗雪": "女",
    "柯孟宜": "女",
    "柯春僖": "女",
    "柯春黛": "女",
    "柯錦鑫": "男",
    "柳書玉": "女",
    "梁庭": "女",
    "楊婉君": "女",
    "楊昌學": "男",
    "楊月娥": "女",
    "楊素粉": "女",
    "楊紫璿": "女",
    "楊美蘭": "女",
    "楊舒涵": "女",
    "楊金珠": "女",
    "楊韻蓉": "女",
    "江佳怡": "女",
    "江盈潔": "女",
    "沈承宗": "男",
    "洪于晴": "女",
    "洪唯": "男",
    "洪瑩蓁": "女",
    "洪筱婷": "女",
    "洪金城": "男",
    "洪雅玲": "女",
    "游建邦": "男",
    "湯玉琦": "女",
    "溫婷伊": "女",
    "潘志滿": "女",
    "潘玟玲": "女",
    "潘秋華": "女",
    "王上苹": "女",
    "王冠蕙": "女",
    "王國美": "女",
    "王均晨": "女",
    "王坤德": "男",
    "王怡誠": "女",
    "王慧玟": "女",
    "王碧宏": "男",
    "王素敏": "女",
    "王芳珠": "女",
    "王韶怡": "女",
    "王馨儀": "女",
    "王麗景": "女",
    "留千惠": "女",
    "白秀清": "女",
    "白豐銘": "男",
    "盧彥妤": "女",
    "盧菀萱": "女",
    "石俊傑": "男",
    "秦育曐": "男",
    "章玲珠": "女",
    "童小鈴": "女",
    "簡達謙": "男",
    "簡麗分": "女",
    "粘振清": "男",
    "粘智淳": "女",
    "紀辰諭": "女",
    "羅怡和": "女",
    "羅賢芳": "男",
    "胡瑜玲": "女",
    "胡詠茹": "女",
    "莊媛婷": "女",
    "莊淑娟": "女",
    "莊郁琳": "女",
    "莊雅惠": "女",
    "葉奕新": "男",
    "葉如玲": "女",
    "葉姿佑": "女",
    "葉桂香": "女",
    "葉瓊雅": "女",
    "葉美鴻": "女",
    "葉迷妮": "女",
    "蔡佩君": "女",
    "蔡宗佑": "男",
    "蔡宗涵": "男",
    "蔡宜均": "女",
    "蔡明真": "女",
    "蔡春美": "女",
    "蔡沛妤": "女",
    "蔡淵輝": "男",
    "蔡石福地": "男",
    "蔡秉諺": "男",
    "蔡純慧": "女",
    "蔡菁雯": "女",
    "蔡雅雅": "女",
    "蔡麗滿": "女",
    "蕭名權": "男",
    "蕭淑勻": "女",
    "蕭淑慧": "女",
    "蕭玉卿": "女",
    "蕭米佋": "女",
    "蕭茲方": "女",
    "薛榕婷": "女",
    "蘇國樑": "男",
    "蘇献珍": "男",
    "蘇相瑜": "女",
    "蘇誌盈": "男",
    "蘇郁合": "女",
    "蘇鴻濱": "男",
    "許婉瑜": "女",
    "許婷婷": "女",
    "許媛婷": "女",
    "許東華": "女",
    "許湘惠": "女",
    "謝佩穎": "女",
    "謝佳容": "女",
    "謝孟君": "女",
    "謝桂红": "女",
    "謝秀枝": "女",
    "謝秀錦": "女",
    "謝耀德": "男",
    "謝萬蒲": "男",
    "賀正楨": "女",
    "賴姳璇": "女",
    "賴幸絹": "女",
    "賴彥帛": "男",
    "賴彩鈴": "女",
    "賴志雄": "男",
    "賴盈達": "男",
    "賴素貞": "女",
    "賴香雪": "女",
    "趙如薰": "女",
    "辛天送": "男",
    "邱品瑄": "女",
    "邱淑春": "女",
    "邵士誠": "男",
    "邵杰": "男",
    "郭宇家": "女",
    "郭宗川": "男",
    "郭富貴": "男",
    "郭瓈灧": "女",
    "郭芳佑": "男",
    "郭芷萱": "女",
    "鄭人豪": "男",
    "鄭伃倢": "女",
    "鄭宇丞": "男",
    "鄭宜涵": "女",
    "鄭斐文": "女",
    "鄭秀娟": "女",
    "鄭美華": "女",
    "鄭翊汝": "女",
    "鄭鈞懿": "女",
    "鄭雅文": "女",
    "鍾昆原": "男",
    "關婉玲": "女",
    "阮善如": "女",
    "阮芬": "女",
    "陳丁清": "男",
    "陳世平": "男",
    "陳冠同": "男",
    "陳君妮": "女",
    "陳品樺": "女",
    "陳威廷": "男",
    "陳定群": "男",
    "陳彥伶": "女",
    "陳慶旺": "男",
    "陳揚旻": "女",
    "陳星妤": "女",
    "陳昭華": "女",
    "陳月美": "女",
    "陳桂宸": "女",
    "陳毓旻": "女",
    "陳淑媛": "女",
    "陳淑芬": "女",
    "陳淑芳": "女",
    "陳珈暐": "男",
    "陳田恬": "女",
    "陳秀美": "女",
    "陳秀花": "女",
    "陳秋芬": "女",
    "陳素葉": "女",
    "陳聖昀": "女",
    "陳芸詞": "女",
    "陳萬和": "男",
    "陳薇莉": "女",
    "陳重仰": "男",
    "陳錦鳳": "女",
    "陳雅惠": "女",
    "陳麗華": "女",
    "韋淑媛": "女",
    "韓倩如": "女",
    "顏妙如": "女",
    "顏宗宏": "男",
    "顏少萱": "女",
    "馮豐隆": "男",
    "高忠偉": "男",
    "高翊綺": "女",
    "高韻筑": "女",
    "魏媛真": "女",
    "黃乙娟": "女",
    "黃佩琪": "女",
    "黃克經": "男",
    "黃品齊": "男",
    "黃壬癸": "男",
    "黃如榛": "女",
    "黃姿蓉": "女",
    "黃忠權": "男",
    "黃惠卿": "女",
    "黃振育": "男",
    "黃政和": "男",
    "黃曉鈺": "女",
    "黃柏豪": "男",
    "黃榮茂": "男",
    "黃沛瑋": "女",
    "黃淑梅": "女",
    "黃淑清": "女",
    "黃淑芳": "女",
    "黃玉宙": "女",
    "黃瓊慧": "女",
    "黃翠月": "女",
    "黃蔚菁": "女",
    "黃金輝": "男",
    "黃麗燕": "女",
}
            gender = GENDER_MAP.get(actual_name, "男")
            timestamp = datetime.datetime.now(pytz.timezone("Asia/Taipei")).strftime("%Y-%m-%d %H:%M:%S")
            new_row = {
                "時間戳記": timestamp,
                "組別": group_selected,
                "保溫組長": leader,
                "個人姓名": actual_name,
                "電子郵件": email,
                "個人編號": auto_member_id,
                "性別": gender,
                "當週回報期間": period,
                "當週填表張數": int(weekly_sheets),
                "累積填表張數": int(cum_sheets),
                "本週組內共學": "是" if co_learn else "否",
                "心得或對法師提問": feedback.strip() if feedback.strip() else "無",
                "最常出現的身體反應": body_reactions.strip() if body_reactions.strip() else "無",
                "最常出現的情緒類別": emotion_types.strip() if emotion_types.strip() else "無",
                "最常用的落地方式": grounding_methods.strip() if grounding_methods.strip() else "無"
            }
            
            success, msg = save_record(new_row, gas_url=gas_api_url)
            if success:
                # Send email copy if actual email was entered
                email_status_msg = ""
                if student_email_val.strip() and "@" in student_email_val:
                    email_success, email_msg = send_email_copy(student_email_val.strip(), new_row)
                    if email_success:
                        email_status_msg = f" 📨 {email_msg}"
                    else:
                        email_status_msg = f" ⚠️ 填報成功，但副本發送失敗：{email_msg}"
                else:
                    email_status_msg = " 💡 (提示：填寫您的真實電子郵件，下次送出即可自動收到填報副本確認信喔！)"
                
                st.session_state["submitted_successfully"] = True
                st.session_state["last_submitted_name"] = actual_name
                st.session_state["last_submitted_msg"] = msg + email_status_msg
                st.session_state["last_submitted_id"] = auto_member_id
                st.session_state["last_submitted_period"] = period
                st.session_state["last_submitted_co_learn"] = "是" if co_learn else "否"
                st.session_state["last_submitted_weekly"] = int(weekly_sheets)
                st.session_state["last_submitted_cum"] = int(cum_sheets)
                st.session_state["last_submitted_body_reactions"] = body_reactions.strip() if body_reactions.strip() else "無"
                st.session_state["last_submitted_emotion_types"] = emotion_types.strip() if emotion_types.strip() else "無"
                st.session_state["last_submitted_grounding_methods"] = grounding_methods.strip() if grounding_methods.strip() else "無"
                st.rerun()

    # Footer login expander for group leaders / admin
    st.write("---")
    with st.expander("🔑 關懷管理登入 (組長 / 總負責人專用)"):
        st.write("請輸入您的管理密碼解鎖權限：")
        login_pwd = st.text_input("輸入登入密碼", type="password", key="footer_pwd_input")
        if st.button("確認登入 🔓", key="footer_login_submit"):
            if login_pwd == admin_password:
                st.session_state["is_authenticated"] = True
                st.session_state["role"] = "📊 總窗口後台管理"
                st.session_state["logged_in_group"] = None
                st.success("🔓 總窗口解鎖成功！您擁有全局最高權限。")
                st.rerun()
            else:
                matched_group = None
                for g, p in leader_passwords.items():
                    if login_pwd == p:
                        matched_group = g
                        break
                if matched_group:
                    st.session_state["is_authenticated"] = True
                    st.session_state["role"] = "👥 保溫組長專區"
                    st.session_state["logged_in_group"] = matched_group
                    st.success(f"🔓 登入成功！您已解鎖為【{matched_group}】權限。")
                    st.rerun()
                else:
                    st.error("密碼錯誤")

elif role == "👥 保溫組長專區":
    if not is_authenticated:
        st.info("👈 請在控制面板輸入「組長密碼」解鎖您的保溫組別權限。")
        st.image("https://img.icons8.com/clouds/400/lock.png", width=250)
    else:
        st.subheader(f"👥 【{logged_in_group}】保溫組長一鍵關懷與心得導出專區")
        st.caption("🔒 系統已安全限制：您目前僅可查閱貴組的填報資料，無法查看其他組別。")
        
        # 1. Primary care selection - select EXACTLY ONE target week
        st.markdown("### 📅 第一步：選擇您要查詢與關懷的「目標週別」")
        reminder_target_week = st.selectbox("請選擇目標週別 *", weeks_list, index=default_week_idx, key="leader_target_week")
        
        target_wk_label = reminder_target_week.split(' (')[0] # e.g. W1
        
        # Load data and build roster list
        all_data = load_data()
        group_roster = groups_data[logged_in_group]["members"]
        g_num = logged_in_group.replace("第 ", "").replace(" 組", "")
        
        roster_list = []
        for idx, m_full in enumerate(group_roster):
            if " " in m_full:
                parts = m_full.split(" ")
                m_id = parts[0]
                m_name = parts[1]
            else:
                m_id = f"{g_num}-{idx+1}"
                m_name = m_full
            roster_list.append({
                "組員號碼": m_id,
                "個人姓名": m_name,
                "編號排序": idx+1
            })
        roster_base_df = pd.DataFrame(roster_list)
        
        # Submissions for selected target week in logged in group
        target_wk_submissions = all_data[(all_data["組別"] == logged_in_group) & (all_data["當週回報期間"] == reminder_target_week)]
        submitted_names = list(set(target_wk_submissions["個人姓名"].tolist()))
        
        # Unsubmitted members
        unsubmitted_members = []
        for row in roster_list:
            if row["個人姓名"] not in submitted_names:
                unsubmitted_members.append(row)
                
        # Simple statistics card
        total_group_m = len(roster_list)
        sub_count = len(submitted_names)
        unsub_count = len(unsubmitted_members)
        sub_rate = (sub_count / total_group_m) * 100 if total_group_m > 0 else 0
        co_learn_count = len(target_wk_submissions[target_wk_submissions["本週組內共學"] == "是"]) if not target_wk_submissions.empty and "本週組內共學" in target_wk_submissions.columns else 0
        
        st.write("---")
        st.markdown("### 📊 本週填報與共學進度統計")
        col_m1, col_m2, col_m3, col_m4, col_m5 = st.columns(5)
        with col_m1:
            st.metric("應填報人數", f"{total_group_m} 人")
        with col_m2:
            st.metric("已完成登記", f"{sub_count} 人")
        with col_m3:
            st.metric("組內共學出席", f"{co_learn_count} 人")
        with col_m4:
            st.metric("尚未登記人數", f"{unsub_count} 人", delta=f"-{unsub_count}" if unsub_count > 0 else None, delta_color="inverse")
        with col_m5:
            st.metric("當週填報率", f"{sub_rate:.1f}%")
            
        st.write("---")
        
        # 2. Roster Matrix Expander (Hidden by default for simplification!)
        with st.expander("📅 展開歷次週別對照矩陣 (跨週對照表)"):
            st.markdown("您可以選擇自訂的回報週別範圍，以查看多週的對照狀態：")
            start_week = st.selectbox("開始週別：", weeks_list, index=0, key="matrix_start_wk")
            end_week = st.selectbox("結束週別：", weeks_list, index=default_week_idx, key="matrix_end_wk")
            
            try:
                start_idx = weeks_list.index(start_week)
                end_idx = weeks_list.index(end_week)
                if start_idx <= end_idx:
                    selected_weeks_range = weeks_list[start_idx:end_idx+1]
                else:
                    selected_weeks_range = weeks_list[end_idx:start_idx+1]
            except Exception:
                selected_weeks_range = weeks_list[0:4]
                
            matrix_data = roster_base_df.copy()
            for wk in selected_weeks_range:
                wk_lbl = wk.split(" (")[0]
                wk_sub_names = all_data[(all_data["組別"] == logged_in_group) & (all_data["當週回報期間"] == wk)]
                status_col = []
                for m_full in group_roster:
                    m_name = m_full.split(" ")[1] if " " in m_full else m_full
                    user_record = wk_sub_names[wk_sub_names["個人姓名"] == m_name]
                    if not user_record.empty:
                        sheets = user_record["當週填表張數"].values[0]
                        status_col.append(f"✅ 已填 ({int(sheets)}張)")
                    else:
                        status_col.append("❌ 未填")
                matrix_data[wk_lbl] = status_col
            
            display_matrix = matrix_data.drop(columns=["編號排序"]).set_index("組員號碼")
            st.table(display_matrix)
            
        st.write("---")
        
        # 3. Unsubmitted care list & LINE template
        st.markdown(f"### 🚨 {target_wk_label} 未填寫關懷專區")
        if unsub_count == 0:
            st.balloons()
            st.success(f"🎉 太棒了！【{logged_in_group}】在 {target_wk_label} 的全體組員皆已完成填報問卷！")
        else:
            col_u1, col_u2 = st.columns([1, 1])
            with col_u1:
                st.warning(f"⚠️ 尚有 **{unsub_count}** 位老師尚未完成本週填寫：")
                unsub_df = pd.DataFrame(unsubmitted_members).drop(columns=["編號排序"]).set_index("組員號碼")
                st.table(unsub_df)
            with col_u2:
                st.markdown("💬 **LINE 溫馨關懷提醒範本：**")
                unsub_mentions = " ".join([f"@{m['個人姓名']}" for m in unsubmitted_members])
                template_text = f"""📢【情緒觀察表填報提醒】

大家好：

這週大家都還好嗎？
是否常常忘記落地，或是多看見一點自己的情緒呢？
每次用功都是向內觀察的真實用功，記得想到上師會很開心我們好好照顧自己～
提醒大家可以找時間做週總結，看見自己完成的每一步都很棒！

👉 填報連結：
https://emot-aw-ktu54azdnwg7qscoaxawew.streamlit.app/#990d8774

📝 尚未完成填報的老師：
{unsub_mentions}"""
                st.text_area("可以直接複製以下文字發到 LINE 群組中：", value=template_text, height=180, key="line_reminder_text_area")

        st.write("---")
        
        # 4. One-Click docx generator for target week (💡 v35 調整：改至數據表格上方，方便組長一秒導出報告)
        st.markdown(f"### 📝 一鍵生成 {target_wk_label} 心得報告 (.docx)")
        doc_bio = generate_docx_report(
            group_name=logged_in_group,
            leader=groups_data[logged_in_group]["leader"],
            week=reminder_target_week,
            submissions=target_wk_submissions,
            roster_list=roster_list
        )
        btn_col1, btn_col2 = st.columns([1, 2])
        with btn_col1:
            st.download_button(
                label="📥 一鍵導出 Word 格式心得報告",
                data=doc_bio,
                file_name=f"{logged_in_group}_情緒觀察每週心得報告_{target_wk_label}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key="download_docx_report_btn"
            )
        with btn_col2:
            st.success("✅ 當週心得報告已生成！點擊左側按鈕即可下載 Word 檔。")
            
        st.write("---")
        
        # 5. Detailed Data for selected target week (💡 v35 調整：數據預覽放置下方，並支援視窗滾動條)
        st.markdown(f"### 📋 {target_wk_label} 組員填報心得與提問清單")
        st.caption("依據組員號碼順序排列，方便閱讀心得與對法師的提問")
        
        leader_df = target_wk_submissions.copy()
        if not leader_df.empty:
            def get_sort_key(id_val):
                try:
                    parts = str(id_val).split("-")
                    return int(parts[-1])
                except Exception:
                    return 999
            leader_df["_sort_key"] = leader_df["個人編號"].apply(get_sort_key)
            leader_df = leader_df.sort_values(by=["_sort_key"]).drop(columns=["_sort_key"])
            display_cols = ["個人編號", "個人姓名"]
            if "本週組內共學" in leader_df.columns:
                display_cols.append("本週組內共學")
            display_cols.extend(["當週填表張數", "累積填表張數"])
            for col_name in ["最常出現的身體反應", "最常出現的情緒類別", "最常用的落地方式"]:
                if col_name in leader_df.columns:
                    display_cols.append(col_name)
            display_cols.extend(["心得或對法師提問", "時間戳記"])
            leader_df_display = leader_df[display_cols].set_index("個人編號")
            st.table(leader_df_display)
        else:
            st.info(f"💡 目前在 {target_wk_label} 期間內，本小組尚無組員填報資料。")



elif role == "📊 總窗口後台管理":
    if not is_authenticated:
        st.info("👈 請在左側控制面板輸入「總窗口管理密碼」解鎖最高後台管理權限。")
        st.image("https://img.icons8.com/clouds/400/security-shield.png", width=250)
    else:
        st.subheader("📊 總窗口全局資料管理庫")
        
        # Load database
        df = load_data()
        
        # 1. Weekly Completion Rate Assessment
        st.markdown("---")
        st.markdown("### 📊 全營 18 組填寫率快速評估")
        
        target_report_week = st.selectbox("請選擇您要彙整的「目標週別」 *", weeks_list, index=default_week_idx)
        
        if not df.empty:
            # Submissions in selected week
            wk_sub = df[df["當週回報期間"] == target_report_week]
            
            rows_stats = []
            for g_name, g_info in groups_data.items():
                m_list_full = g_info["members"]
                m_list_names = [m.split(" ")[1] if " " in m else m for m in m_list_full]
                ldr = g_info["leader"]
                g_sub = wk_sub[wk_sub["組別"] == g_name] if not wk_sub.empty else pd.DataFrame()
                sub_names = g_sub["個人姓名"].tolist() if not g_sub.empty else []
                sub_count = len(set(sub_names) & set(m_list_names))
                unsub_count = len(m_list_full) - sub_count
                g_co_learn = len(g_sub[g_sub["本週組內共學"] == "是"]) if not g_sub.empty and "本週組內共學" in g_sub.columns else 0
                rate = (sub_count / len(m_list_full)) * 100 if len(m_list_full) > 0 else 0
                g_sheets = int(g_sub["當週填表張數"].sum()) if not g_sub.empty else 0
                avg_sheets = (g_sheets / len(m_list_full)) if len(m_list_full) > 0 else 0.0
                rows_stats.append({
                    "組別": g_name,
                    "保溫組長": ldr,
                    "應填人數": len(m_list_full),
                    "已填人數": sub_count,
                    "未填人數": unsub_count,
                    "共學人數": f"{g_co_learn} 人",
                    "當週填表張數": f"{g_sheets} 張",
                    "平均每人填表數": f"{avg_sheets:.1f} 張/人",
                    "填報率": f"{rate:.1f}%"
                })
            
            stats_df = pd.DataFrame(rows_stats)
            st.table(stats_df[["組別", "保溫組長", "應填人數", "已填人數", "未填人數", "共學人數", "當週填表張數", "平均每人填表數", "填報率"]].set_index("組別"))
            
            # 💡 v45 升級：總窗口一鍵導出全營總報告 (.docx) 與 打包生成 18 組個別組長報告 (.zip)
            target_wk_label = target_report_week.split(' (')[0]
            st.write("---")
            st.markdown(f"### 📝 總窗口報告一鍵導出與打包專區【{target_wk_label}】")
            st.caption("支援一鍵導出包含全營 18 組評估與心得條列之『全營總報告』，或直接打包生成 18 份各組獨立之『組長 Word 報告 ZIP 檔』。")
            
            doc_all_bio = generate_all_groups_docx_report(
                week=target_report_week,
                all_data=df,
                groups_data=groups_data
            )
            
            zip_18_bio = generate_18_groups_zip_reports(
                week=target_report_week,
                all_data=df,
                groups_data=groups_data
            )
            
            btn_col_a1, btn_col_a2 = st.columns([1, 1])
            with btn_col_a1:
                st.download_button(
                    label="📥 導出全營 18 組總總結報告 (.docx)",
                    data=doc_all_bio,
                    file_name=f"園區情緒觀察營_全營18組總心得報告_{target_wk_label}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key="download_all_groups_docx_report_btn"
                )
                st.caption("💡 內含 18 組填寫率評估表與各組組員心得條列摘要。")
                
            with btn_col_a2:
                st.download_button(
                    label="📦 打包下載 18 組個別組長報告 (.zip)",
                    data=zip_18_bio,
                    file_name=f"園區情緒觀察營_18組個別組長報告_{target_wk_label}.zip",
                    mime="application/zip",
                    key="download_18_groups_zip_report_btn"
                )
                st.caption("💡 內含 18 份獨立 Word 檔，方便分發給各組組長。")
        else:
            st.warning("⚠️ 尚無任何數據，無法生成填寫率評估。")
        
        # 2. Database View
        st.markdown("---")
        st.markdown("### 📁 完整資料庫檢視與全局數據篩選器")
        
        if not df.empty:
            total_records = len(df)
            total_weekly_sheets = df["當週填表張數"].sum()
            unique_members = df["個人姓名"].nunique()
            
            m1, m2, m3 = st.columns(3)
            with m1:
                st.metric("營隊總回報筆數", f"{total_records} 筆")
            with m2:
                st.metric("累計當週總填表張數", f"{total_weekly_sheets} 張")
            with m3:
                st.metric("全營不重複填報人數", f"{unique_members} 人")
            
            st.write("---")
            st.markdown("##### 🔍 系統全局數據篩選器")
            f_col1, f_col2 = st.columns(2)
            with f_col1:
                filter_group = st.multiselect("按組別篩選：", options=list(df["組別"].unique()))
            with f_col2:
                filter_week = st.multiselect("按回報期間篩選：", options=list(df["當週回報期間"].unique()))
                
            filtered_df = df.copy()
            if filter_group:
                filtered_df = filtered_df[filtered_df["組別"].isin(filter_group)]
            if filter_week:
                filtered_df = filtered_df[filtered_df["當週回報期間"].isin(filter_week)]
                
            st.write(f"📂 顯示篩選結果：共 {len(filtered_df)} 筆紀錄")
            
            # 💡 v35 調整：改至數據預覽表格上方，方便總窗口一秒匯出，且下方表格已加上視窗滾動條防拉伸
            csv_data = filtered_df.to_csv(index=False, encoding='utf-8-sig')
            st.download_button(
                label="📥 匯出並下載全局 CSV 報表 (可用 Excel 直接開啟)",
                data=csv_data,
                file_name=f"all_groups_emotional_reports_{datetime.datetime.now(pytz.timezone('Asia/Taipei')).strftime('%Y%m%d_%H%M%S')}.csv",
                mime="text/csv"
            )
            
            st.write("") # 增加適度間距
            
            # 💡 v38 調整：將系統全局篩選數據表渲染為 HTML 並套用專屬的 .scroll-table-container 類，維持視窗滾動且表頭凍結
            html_table = filtered_df.set_index("時間戳記").to_html(escape=False, border=0)
            st.markdown(f'''
            <div class="scroll-table-container">
                {html_table}
            </div>
            ''', unsafe_allow_html=True)
        else:
            st.warning("⚠️ 目前資料庫尚無任何填報紀錄。")
